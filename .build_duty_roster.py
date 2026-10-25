import os, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL

# Duty Roster: the Duty Planner's master plan of who is on duty each day, for two duties:
# Base Duty Officer (BDO) on the "BDO Plan" sheet and Squadron Duty Orderly (SDO) on the "SDO Plan" sheet.
# Each day holds the department on that duty; the Duty Crew sheet lists who can do each duty (the departments pick the person).
# Both calendars use exactly the same cell layout as the department planners' SDO Plan, so the dashboard can copy days across.
OUT = sys.argv[1]
YEAR = 2026
ND = 15
NDS = 10
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTHS_FULL = ["January", "February", "March", "April", "May", "June", "July", "August",
               "September", "October", "November", "December"]
FONT = "Arial"
DEPTS = ["SUP 1", "SUP 2", "Loadmaster", "TACCO", "Pilot"]
NC = 50   # rows per duty crew list
# sample crew lists (callsign, name, department): SDO from SUP 1, SUP 2 and Loadmaster; BDO officers from any department
import seedplan
SDO_CREW, BDO_CREW = seedplan.SDO_CREW, seedplan.BDO_CREW
DUTIES = [("BDO", "Base Duty Officer", "BDO Plan", BDO_CREW, 1, "1D4ED8"), ("SDO", "Squadron Duty Orderly", "SDO Plan", SDO_CREW, 5, "BE123C")]   # code, name, sheet, sample crew, first column on Duty Crew, colour
DEPT_COL = [("DBEAFE", "1E3A8A"), ("FECDD3", "881337"), ("D1FAE5", "065F46"), ("FEF3C7", "78350F"),
            ("EDE9FE", "5B21B6"), ("CFFAFE", "164E63"), ("FCE7F3", "9D174D"), ("E0E7FF", "3730A3")]
SLATE900, SLATE800, SLATE700, SLATE500, SLATE400 = "0F172A", "1E293B", "334155", "64748B", "94A3B8"
SLATE200, SLATE100 = "E2E8F0", "F1F5F9"
AMBER300 = "FCD34D"

def f(bold=False, color="0F172A", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="CBD5E1")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
INPUT_FILL = fill("FFFBEB")

wb = Workbook()
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
def banner(ws, title, subtitle, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    for r in (1, 2):
        for c in range(1, last_col + 1):
            ws.cell(r, c).fill = fill(SLATE900)
    ws.cell(1, 1, title).font = f(True, "FFFFFF", 16)
    ws.cell(2, 1, subtitle).font = f(False, SLATE400, 10)
    ws.cell(1, 1).alignment = LEFT; ws.cell(2, 1).alignment = LEFT
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 18
def header(ws, row, col, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col + i, lab)
        c.font = f(True, "FFFFFF", 10); c.fill = fill(SLATE900); c.alignment = CENTER; c.border = BOX
def apply_dept_cf(ws, rng, tl, duty=None):
    for k, (bg, fg) in enumerate(DEPT_COL):
        val = f'IFERROR(INDEX({duty}Dept,MATCH({tl},{duty}Call,0)),"")' if duty else tl
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND({tl}<>"",{val}=Departments!$A${6 + k})'], fill=fill(bg), font=Font(color=fg, bold=True), stopIfTrue=True))

dash = wb.active; dash.title = "Dashboard"
plans = {"BDO": wb.create_sheet("BDO Plan"), "SDO": wb.create_sheet("SDO Plan")}
crew = wb.create_sheet("Duty Crew")
deps = wb.create_sheet("Departments")

# ---------------------------------------------------------------- Dashboard (year + how-to)
banner(dash, "DUTY ROSTER  |  BDO AND SDO", "Kept by the Duty Planner. Base Duty Officer (BDO) and Squadron Duty Orderly (SDO): a person from each duty's crew list on every day.", 8)
dash["A4"] = "YEAR"; dash["A4"].font = f(True, SLATE500, 9)
dash["A5"] = YEAR; dash["A5"].font = f(True, "1D4ED8", 14); dash["A5"].fill = INPUT_FILL; dash["A5"].border = BOX; dash["A5"].alignment = CENTER
dash["G4"] = "FILE"; dash["G4"].font = f(True, SLATE500, 9)
dash["G5"] = "Duty Roster"; dash["G5"].font = f(True, "BE123C", 14)
dv = DataValidation(type="list", formula1='"' + ",".join(str(y) for y in range(2025, 2036)) + '"'); dv.add("A5"); dash.add_data_validation(dv)
name("PlanYear", "Dashboard!$A$5")
steps = [
    ("HOW DUTY PLANNING WORKS", True),
    ("1. Duty crew: on the Duty Crew sheet, list who can do each duty (callsign, name, department). BDO and SDO have separate lists.", False),
    ("2. Duty Planner: open this file in the planner dashboard (Open my Excel), then Add other departments to see every department's plans read-only.", False),
    ("3. On the dashboard's Duty Planner tab, pick the department for BDO and for SDO on each day. Save to Excel writes only to this file.", False),
    ("4. Each department: open your own planner in the dashboard with this Duty Roster alongside. The Duty Planner tab shows what differs; click 'Copy my duty days from the Duty Roster', then Save to Excel.", False),
    ("5. Then put someone from the crew list on BDO / SDO for each of your duty days on the month tabs. The DUTY DAYS row turns red on a day with nobody on that duty.", False),
    ("", False),
    ("You can also plan straight in Excel on the BDO Plan and SDO Plan sheets: pick a department from the dropdown in each day. The table on the right lists each department's dates.", False),
    ("Departments are listed on the Departments sheet. Use the same spelling as the Department name in each department's planner (its Dashboard, cell G5).", False),
]
for i, (t, h) in enumerate(steps):
    dash.merge_cells(start_row=8 + i, start_column=1, end_row=8 + i, end_column=8)
    c = dash.cell(8 + i, 1, t); c.font = f(True, SLATE900, 12) if h else f(False, SLATE700, 11)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    dash.row_dimensions[8 + i].height = 22 if h or not t else 32
for col, w in zip("ABCDEFGH", [12, 14, 14, 14, 14, 14, 16, 14]):
    dash.column_dimensions[col].width = w

# ---------------------------------------------------------------- Departments
banner(deps, "DEPARTMENTS", "Add a department in the next empty cell (no gaps). Spell it exactly as in that department's planner (Dashboard, cell G5).", 3)
header(deps, 5, 1, ["Department"])
for k in range(ND):
    c = deps.cell(6 + k, 1, DEPTS[k] if k < len(DEPTS) else None); c.border = BOX
    if k < len(DEPT_COL):
        c.fill = fill(DEPT_COL[k][0]); c.font = f(True, DEPT_COL[k][1])
    else:
        c.fill = INPUT_FILL; c.font = f(True)
deps.column_dimensions["A"].width = 22
name("Depts", f"Departments!$A$6:INDEX(Departments!$A$6:$A${5+ND},MAX(1,COUNTA(Departments!$A$6:$A${5+ND})))")

# ---------------------------------------------------------------- Duty Crew: who can do each duty
banner(crew, "DUTY CREW", "Who can do each duty. Add people in the empty rows (no gaps). Department must match the Departments sheet.", 7)
for code, nm, sh, sample, c0, col in DUTIES:
    crew.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
    t = crew.cell(4, c0, f"{code}  |  {nm.upper()}"); t.font = f(True, "FFFFFF", 11); t.fill = fill(col); t.alignment = CENTER
    for j in (1, 2): crew.cell(4, c0 + j).fill = fill(col)
    header(crew, 5, c0, ["Callsign", "Name", "Department"])
    for i in range(NC):
        r = 6 + i
        for j in range(3):
            cell = crew.cell(r, c0 + j); cell.fill = INPUT_FILL; cell.border = BOX; cell.font = f(j == 0, "1D4ED8")
        if i < len(sample):
            for j, v in enumerate(sample[i]): crew.cell(r, c0 + j, v)
    dv = DataValidation(type="list", formula1="Depts", allow_blank=True); dv.add(f"{CL(c0 + 2)}6:{CL(c0 + 2)}{5 + NC}"); crew.add_data_validation(dv)
    apply_dept_cf(crew, f"{CL(c0 + 2)}6:{CL(c0 + 2)}{5 + NC}", f"{CL(c0 + 2)}6")
    for j, w in enumerate([12, 26, 14]): crew.column_dimensions[CL(c0 + j)].width = w
    name(f"{code}Call", f"'Duty Crew'!${CL(c0)}$6:${CL(c0)}${5 + NC}"); name(f"{code}Dept", f"'Duty Crew'!${CL(c0 + 2)}$6:${CL(c0 + 2)}${5 + NC}")
crew.column_dimensions["D"].width = 3
crew.freeze_panes = "A6"
crew.sheet_properties.tabColor = "7C3AED"

# ---------------------------------------------------------------- BDO Plan and SDO Plan (same layout as the department planners' SDO Plan)
SDO_BW, SDO_BH = 8, 16
def sdo_block(mi):
    return 1 + (mi % 3) * SDO_BW, 6 + (mi // 3) * SDO_BH
SDO_LAST = 3 * SDO_BW - 1
SC0 = SDO_LAST + 3
DOW = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
for code, nm, shname, sample, cc0, col in DUTIES:
    sdo = plans[code]
    banner(sdo, f"{code} PLANNING CALENDAR  |  {nm.upper()}", f"Pick the department on {code} duty for each day. The table on the right lists each department's dates. Who can do {code}: the Duty Crew sheet.", SC0 + NDS)
    sdo["A4"] = "HIGHLIGHT DEPT"; sdo["A4"].font = f(True, SLATE500, 9)
    sdo.merge_cells("A4:B4"); sdo.merge_cells("C4:D4")
    sdo["C4"].fill = INPUT_FILL; sdo["C4"].border = BOX; sdo["C4"].font = f(True, "1D4ED8", 11); sdo["C4"].alignment = CENTER
    sdo["E4"] = "Pick a department to highlight its days."; sdo["E4"].font = f(italic=True, color=SLATE500, size=9)
    dv = DataValidation(type="list", formula1="Depts", allow_blank=True); dv.add("C4"); sdo.add_data_validation(dv)
    sdo.row_dimensions[4].height = 22
    dv_sdo = DataValidation(type="list", formula1="Depts", allow_blank=True, showErrorMessage=True, errorStyle="stop",
                            errorTitle=f"{code} department", error="Pick a department from the list. Departments are listed on the Departments sheet.")
    blocks = {}
    for mi in range(13):   # 12 months + next January (the look-ahead block sits under January, rows 70-85)
        c0, r0 = sdo_block(mi)
        first = f"DATE(PlanYear,{mi + 1},1)"
        blocks[mi] = f"${CL(c0)}${r0 + 3}:${CL(c0 + 6)}${r0 + 13}"
        sdo.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 6)
        t = sdo.cell(r0, c0, f'=UPPER(TEXT({first},"mmmm yyyy"))&"  ·  {code}"'); t.font = f(True, "FFFFFF", 11); t.alignment = CENTER
        for c in range(c0, c0 + 7):
            sdo.cell(r0, c).fill = fill(col if c == c0 else SLATE900)
            h = sdo.cell(r0 + 1, c, DOW[c - c0]); h.font = f(True, AMBER300 if c - c0 >= 5 else "FFFFFF", 9); h.fill = fill(SLATE700 if c - c0 >= 5 else SLATE800); h.alignment = CENTER
        for c in range(c0, c0 + 7): sdo.cell(r0, c).fill = fill(col)
        sdo.row_dimensions[r0].height = 20; sdo.row_dimensions[r0 + 1].height = 15
        for w in range(6):
            rd, rp = r0 + 2 + 2 * w, r0 + 3 + 2 * w
            for wd in range(7):
                c = c0 + wd; n = w * 7 + wd + 1
                dcell = sdo.cell(rd, c, f'=IF(AND({n}-(WEEKDAY({first},2)-1)>=1,{n}-(WEEKDAY({first},2)-1)<=DAY(EOMONTH({first},0))),{first}+{n}-WEEKDAY({first},2),"")')
                dcell.number_format = "d"; dcell.font = f(True, SLATE500, 9); dcell.alignment = Alignment(horizontal="left", vertical="top")
                dcell.border = Border(left=thin, right=thin, top=thin)
                pcell = sdo.cell(rp, c); pcell.fill = INPUT_FILL; pcell.font = f(True, "1D4ED8", 10)
                pcell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True); pcell.border = Border(left=thin, right=thin, bottom=thin)
            sdo.row_dimensions[rd].height = 14; sdo.row_dimensions[rp].height = 24
            L0, L6 = CL(c0), CL(c0 + 6)
            rng = f"{L0}{rp}:{L6}{rp}"
            dv_sdo.add(rng)
            sdo.conditional_formatting.add(f"{L0}{rd}:{L6}{rp}", FormulaRule(formula=[f'{L0}${rd}=""'], fill=fill(SLATE200), stopIfTrue=True))
            sdo.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($C$4<>"",{L0}{rp}=$C$4)'], fill=fill(SLATE900), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
            apply_dept_cf(sdo, rng, f"{L0}{rp}")
        for c in range(c0, c0 + 7):
            sdo.column_dimensions[CL(c)].width = 12
        if mi % 3 < 2:
            sdo.column_dimensions[CL(c0 + 7)].width = 2
        if mi % 3 == 0:
            sdo.row_dimensions[r0 + SDO_BH - 2].height = 10; sdo.row_dimensions[r0 + SDO_BH - 1].height = 10
    sdo.add_data_validation(dv_sdo)
    sdo.column_dimensions[CL(SDO_LAST + 1)].width = 2; sdo.column_dimensions[CL(SDO_LAST + 2)].width = 2
    sdo.merge_cells(start_row=4, start_column=SC0, end_row=4, end_column=SC0 + NDS + 1)
    sdo.cell(4, SC0, f"{code} DATES BY DEPARTMENT (filled in automatically)").font = f(True, SLATE900, 11)
    header(sdo, 5, SC0, ["Month"])
    sdo.column_dimensions[CL(SC0)].width = 11
    for k in range(NDS):
        cc = SC0 + 1 + k; L = CL(cc)
        h = sdo.cell(5, cc, f'=IF(Departments!$A${6 + k}="","",Departments!$A${6 + k})'); h.font = f(True, "FFFFFF", 10); h.fill = fill(SLATE900); h.alignment = CENTER; h.border = BOX
        sdo.column_dimensions[L].width = 15
        for mi in range(13):
            r = 6 + mi + (mi == 12); first = f"DATE(PlanYear,{mi + 1},1)"
            terms = "&".join(f'IF(INDEX({blocks[mi]},2*INT(({d}+WEEKDAY({first},2)-1)/7)+1,MOD({d}+WEEKDAY({first},2)-1,7)+1)={L}$5,", "&{d + 1},"")' for d in range(31))
            c = sdo.cell(r, cc, f'=IF({L}$5="","",MID({terms},3,200))')
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True); c.font = f(False, SLATE800, 9); c.border = BOX
        cnt = "+".join(f"COUNTIF({blocks[mi]},{L}$5)" for mi in range(12))
        c = sdo.cell(18, cc, f'=IF({L}$5="","",{cnt})'); c.font = f(True); c.alignment = CENTER; c.border = BOX; c.fill = fill(SLATE100)
    gcol = SC0 + 1 + NDS
    header(sdo, 5, gcol, ["Days not set"]); sdo.column_dimensions[CL(gcol)].width = 12
    for mi in range(13):
        r = 6 + mi + (mi == 12)   # next January goes under the total
        sdo.cell(r, SC0, MONTHS_FULL[mi] if mi < 12 else "Jan (next yr)").font = f(True, SLATE800, 9); sdo.cell(r, SC0).border = BOX; sdo.cell(r, SC0).alignment = Alignment(vertical="top")
        g = sdo.cell(r, gcol, f"=DAY(EOMONTH(DATE(PlanYear,{mi + 1},1),0))-COUNTIF({blocks[mi]},\"?*\")"); g.alignment = CENTER; g.border = BOX; g.font = f(True, "B91C1C", 10); g.number_format = "0;-0;-"
        sdo.row_dimensions[r].height = max(sdo.row_dimensions[r].height or 15, 26)
    sdo.cell(18, SC0, "Total days").font = f(True); sdo.cell(18, SC0).fill = fill(SLATE100); sdo.cell(18, SC0).border = BOX
    apply_dept_cf(sdo, f"{CL(SC0 + 1)}5:{CL(SC0 + NDS)}5", f"{CL(SC0 + 1)}5")
    sdo.freeze_panes = "A6"
    sdo.sheet_properties.tabColor = col
    # sample: October (SDO matches the department samples except 9 Oct goes to Loadmaster, to show a difference to copy)
    for dd, dep in seedplan.DUTY[code].items():
        c0, r0 = sdo_block(dd.month - 1); n = dd.day - 1 + dt.date(YEAR, dd.month, 1).weekday()
        sdo.cell(r0 + 3 + 2 * (n // 7), c0 + n % 7, dep)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.active = 0
# layout stamp: the dashboard compares it with the layout it was written for (bump it when cell positions change)
dash["Z2"] = "LODESTAR roster layout 2"; dash.column_dimensions["Z"].hidden = True
wb.save(OUT)
print("saved", OUT)
