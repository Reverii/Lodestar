import sys, calendar, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL

# Master Planner: the master list of planned flights for every department.
# Flights sheet: one flight per row, rows 6..1005, columns A Date, B Shift (AM / PM / NIGHT), C Flight code, D Time, E Crew (override),
#   F Comment, G Sorties (formula), H Crew used (formula). Any number of flights per shift; order within a shift = row order.
# Crew: departments the flight needs, as tags P (Pilot), L (Loadmaster), T (TACCO), S1 (SUP 1), S2 (SUP 2), e.g. "P S2".
#   Blank = the flight code's crew (Flight Codes column E); blank there too = every department.
# Day Status sheet: one row per date of the year from row 6 (A date, B day, C status). Blank = flying day; a reason = no flying.
# Month tabs: a read-only calendar built from the two lists by formulas. The planner dashboard reads and writes the lists, not the calendar.
OUT = sys.argv[1]
YEAR = 2026
FONT = "Arial"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTHS13 = MONTHS + ["Jan (next year)"]   # look-ahead: January of the next year, planned while still in December
MONTHS_FULL = list(calendar.month_name)[1:]
SHIFTS = ["AM", "PM", "NIGHT"]
SHIFT_STYLE = [("F0F9FF", "075985"), ("FFFBEB", "92400E"), ("1E1B4B", "C7D2FE")]   # AM, PM, NIGHT: fill, font
SLATE900, SLATE700, SLATE500, SLATE400, SLATE300, SLATE200, SLATE100 = "0F172A", "334155", "64748B", "94A3B8", "CBD5E1", "E2E8F0", "F1F5F9"
NF = 1000          # flight rows 6..1005
NC = 40            # flight code rows on the Flight Codes sheet
LINES = 3          # flights shown per shift in the Excel calendar (more show as "+n more")
WEEK_ROWS = 2 + 3 * LINES   # header, day status, 3 shifts x LINES

def f(bold=False, color=SLATE900, size=10, italic=False): return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color=SLATE300); hair = Side(style="thin", color=SLATE200); heavy = Side(style="medium", color=SLATE700)
BOX = Border(left=hair, right=hair, top=hair, bottom=hair)
CENTER = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)

wb = Workbook()
dash = wb.active; dash.title = "Dashboard"
fl = wb.create_sheet("Flights")
days = wb.create_sheet("Day Status")
codes = wb.create_sheet("Flight Codes")

def banner(ws, title, sub, last):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
        for c in range(1, last + 1): ws.cell(r, c).fill = fill(SLATE900)
    ws.cell(1, 1, title).font = f(True, "FFFFFF", 16); ws.cell(2, 1, sub).font = f(False, SLATE400, 10)
    ws.cell(1, 1).alignment = LEFT; ws.cell(2, 1).alignment = LEFT
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 18

def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
name("PlanYear", "Dashboard!$A$5")   # the planning year: every date and calendar in this file follows it
def header(ws, row, heads, colors=None):
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row, i, h); c.font = f(True, "FFFFFF"); c.fill = fill((colors or {}).get(i, SLATE900)); c.alignment = CENTER; c.border = BOX
    ws.row_dimensions[row].height = 22

# ---- Flight Codes (crew per code) and the day status reasons
banner(codes, "FLIGHT CODES", "The flights programme: every flight code the squadron uses. Departments copy this list into their Activities sheet in the planner dashboard. Crew = the departments that fly this code (blank = every department). Mission Type is used by the currency rules.", 7)
header(codes, 5, ["Code", "Description", "Sorties", "Flight Hrs", "Crew", "Mission Type"], {5: "7E22CE", 6: "5B21B6"})
CATALOG = [("FLT-1.5", "Routine Sortie (1.5h)", 1, 1.5, ""), ("FLT-2.0", "Standard Sortie (2.0h)", 1, 2.0, ""), ("FLT-3.0", "Extended Mission (3.0h)", 1, 3.0, ""),
           ("NVG-2.0", "Night Flying Sortie (2.0h)", 1, 2.0, ""), ("2xFLT-1.5", "Two routine sorties in one shift", 2, 3.0, ""), ("2xFLT-2.0", "Two standard sorties in one shift", 2, 4.0, ""),
           ("MSN-3.0", "Mission sortie: Pilot, TACCO, SUP 2 (3.0h)", 1, 3.0, "P T S2"), ("PLT-TRG", "Pilot training sortie (1.5h)", 1, 1.5, "P"),
           ("EIS", "EIS training sortie (2.0h)", 1, 2.0, "S2"), ("MAS TRG", "MAS tactical training (3.0h)", 1, 3.0, "S2")]
SYLTAG = {"EIS": "EIS", "MAS TRG": "MAS Tact"}   # Flight Codes H: the cell picker lists only syllabus items whose Track or Description contains it
MTYPE = {"FLT-1.5": "GEN", "FLT-2.0": "GEN", "FLT-3.0": "GEN", "NVG-2.0": "NIGHT", "2xFLT-1.5": "GEN", "2xFLT-2.0": "GEN", "MSN-3.0": "MSN", "PLT-TRG": "TRG", "EIS": "TRG", "MAS TRG": "TRG"}
DAY_STATUS = ["No Flying", "Stand-down", "Maintenance", "Public Holiday"]
import realdata
REAL = realdata.ON   # real squadron data from LODESTAR_Data_Input.xlsx
if REAL:
    CATALOG = [(c["code"], c["desc"], c["sorties"], c["hrs"], c["crew"]) for c in realdata.FLIGHT_CODES]
    MTYPE = {c["code"]: c["mtype"] or None for c in realdata.FLIGHT_CODES}
    SYLTAG = {c["code"]: c["tag"] or None for c in realdata.FLIGHT_CODES}
    DAY_STATUS = realdata.DAY_STATUS[:12] or DAY_STATUS
for i in range(NC):
    r = 6 + i
    for c in range(1, 7):
        codes.cell(r, c).border = BOX; codes.cell(r, c).font = f(color="1D4ED8"); codes.cell(r, c).fill = fill("FFFBEB")
    if i < len(CATALOG):
        for j, v in enumerate(CATALOG[i]): codes.cell(r, 1 + j, v)
        codes.cell(r, 6, MTYPE.get(CATALOG[i][0])); codes.cell(r, 8, SYLTAG.get(CATALOG[i][0]))
    cell = codes.cell(r, 8); cell.border = BOX; cell.font = f(color="92400E"); cell.fill = fill("FEF3C7")
for col, w in zip("ABCDEFG", [14, 40, 10, 11, 12, 14, 26]): codes.column_dimensions[col].width = w
c = codes.cell(5, 7, "Day status (no flying)"); c.font = f(True, "FFFFFF"); c.fill = fill("B91C1C"); c.alignment = CENTER
for i in range(12):
    cell = codes.cell(6 + i, 7, DAY_STATUS[i] if i < len(DAY_STATUS) else None); cell.fill = fill("FEF2F2"); cell.font = f(color="B91C1C"); cell.border = BOX
codes.cell(19, 7, "Add your own reasons in the empty rows.").font = f(False, SLATE500, 9, True)
GROUND_LIST = ["Ground School", "Safety Brief", "CRM Lecture", "Sim Training", "Sqn Meeting", "Admin Day", "Fitness / IPPT", "Medical"]
if REAL and realdata.GROUND_LIST: GROUND_LIST = realdata.GROUND_LIST[:20]
c = codes.cell(5, 9, "Ground programme"); c.font = f(True, "FFFFFF"); c.fill = fill("15803D"); c.alignment = CENTER
for i in range(20):
    cell = codes.cell(6 + i, 9, GROUND_LIST[i] if i < len(GROUND_LIST) else None); cell.fill = fill("F0FDF4"); cell.font = f(color="166534"); cell.border = BOX
codes.cell(26, 9, "The ground programme list (any text is allowed too).").font = f(False, SLATE500, 9, True)
codes.column_dimensions["H"].width = 15; codes.column_dimensions["I"].width = 24
c = codes.cell(5, 8, "Syllabus tag"); c.font = f(True, "FFFFFF"); c.fill = fill("B45309"); c.alignment = CENTER
codes.cell(4, 8, "e.g. EIS, MAS Tact").font = f(False, SLATE500, 8, True)
name("GrList", "OFFSET('Flight Codes'!$I$6,0,0,MAX(1,COUNTA('Flight Codes'!$I$6:$I$25)),1)")
dv = DataValidation(allow_blank=True, showInputMessage=True, promptTitle="Crew", prompt="Departments that fly this code: P Pilot, L Loadmaster, T TACCO, S1 SUP 1, S2 SUP 2. Blank = every department.")
dv.add(f"E6:E{5 + NC}"); codes.add_data_validation(dv)
codes.freeze_panes = "A6"
name("FlCodes", f"OFFSET('Flight Codes'!$A$6,0,0,MAX(1,COUNTA('Flight Codes'!$A$6:$A${5 + NC})),1)")
name("CodeTable", f"'Flight Codes'!$A$6:$E${5 + NC}")
name("DayStatusList", "OFFSET('Flight Codes'!$G$6,0,0,MAX(1,COUNTA('Flight Codes'!$G$6:$G$17)),1)")

# ---- Flights: the list
banner(fl, "FLIGHTS", "The master list: one flight per row, any number per shift. Add flights in the empty rows (or with + Add flight in the planner dashboard). Crew blank = the code's crew.", 8)
fl["A3"] = "Sort or filter as you like. To remove a flight, clear its row. Grey columns fill in by themselves."; fl["A3"].font = f(False, SLATE500, 9, True)
header(fl, 5, ["Date", "Shift", "Flight code", "Time", "Crew (only if different)", "Comment", "Sorties", "Crew used", "Key", "Order", "Counts", "Hrs"], {5: "7E22CE", 7: SLATE500, 8: SLATE500, 9: SLATE400, 10: SLATE400, 11: SLATE400, 12: SLATE400})
import seedplan
rows = [f[:4] + (f[4], f[5]) for f in seedplan.FLIGHTS]   # (date, shift, code, time, comment, crew)
for i in range(NF):
    r = 6 + i
    for c in range(1, 13):
        cell = fl.cell(r, c); cell.border = BOX
        cell.fill = fill(SLATE100) if c >= 7 else fill("FFFBEB"); cell.font = f(c == 3, SLATE500 if c >= 7 else "1D4ED8", 10)
    fl.cell(r, 1).number_format = "ddd dd mmm yyyy"; fl.cell(r, 4).number_format = "@"; fl.cell(r, 2).alignment = CENTER; fl.cell(r, 3).alignment = CENTER
    fl.cell(r, 7, f'=IF(C{r}="","",IFERROR(VLOOKUP(C{r},CodeTable,3,0),1))').alignment = CENTER
    fl.cell(r, 8, f'=IF(C{r}="","",IF(E{r}<>"",E{r},IF(IFERROR(VLOOKUP(C{r},CodeTable,5,0)&"","")="","All",VLOOKUP(C{r},CodeTable,5,0)&"")))').alignment = CENTER
    # flights in a shift are numbered by time (blank time last), then by row: key = date|shift|n
    fl.cell(r, 10, f'=IF(C{r}="","",IF(D{r}="","zz","t"&D{r}))').font = f(False, SLATE400, 8)
    # Counts: 1 for a flight on a flying day (0 on a no-flying day); Hrs: its planned flying hours from the Flight Codes sheet
    fl.cell(r, 11, f'=IF(OR(A{r}="",C{r}=""),"",IF(IFERROR(INDEX(DayStat,A{r}-DATE(PlanYear,1,1)+1)&"","")<>"",0,1))').font = f(False, SLATE400, 8)
    fl.cell(r, 12, f'=IF(K{r}="","",K{r}*IFERROR(VLOOKUP(C{r},CodeTable,4,0),0))').font = f(False, SLATE400, 8)
    fl.cell(r, 9, f'=IF(OR(A{r}="",C{r}=""),"",A{r}&"|"&B{r}&"|"&(COUNTIFS(FlDate,A{r},FlShift,B{r},FlCode,"<>",FlOrder,"<"&J{r})+COUNTIFS(A$6:A{r},A{r},B$6:B{r},B{r},C$6:C{r},"<>",J$6:J{r},J{r})))').font = f(False, SLATE400, 8)
    if i < len(rows):
        dte, sh, cd, tm, cm, cw = rows[i]
        fl.cell(r, 1, dte); fl.cell(r, 2, sh); fl.cell(r, 3, cd); fl.cell(r, 4, tm)
        if cw: fl.cell(r, 5, cw)
        if cm: fl.cell(r, 6, cm)
for col, w in zip("ABCDEFGHIJKL", [17, 9, 13, 13, 20, 30, 9, 11, 14, 10, 8, 8]): fl.column_dimensions[col].width = w
fl.freeze_panes = "A6"; fl.auto_filter.ref = f"A5:L{5 + NF}"
dv = DataValidation(type="list", formula1='"AM,PM,NIGHT"', allow_blank=True); dv.add(f"B6:B{5 + NF}"); fl.add_data_validation(dv)
dv = DataValidation(type="list", formula1="FlCodes", allow_blank=True, showErrorMessage=True, errorStyle="warning", errorTitle="Unknown flight code", error="That code isn't on the Flight Codes sheet. Keep it anyway?")
dv.add(f"C6:C{5 + NF}"); fl.add_data_validation(dv)
dv = DataValidation(type="date", operator="between", formula1="DATE(PlanYear,1,1)", formula2="DATE(PlanYear+1,1,31)", allow_blank=True, showErrorMessage=True, errorTitle="Date", error="Use a date in the planning year (Dashboard A5) or January of the next year.")
dv.add(f"A6:A{5 + NF}"); fl.add_data_validation(dv)
dv = DataValidation(allow_blank=True, showInputMessage=True, promptTitle="Crew", prompt="Only when this flight differs from its code: P Pilot, L Loadmaster, T TACCO, S1 SUP 1, S2 SUP 2 (e.g. P S2).")
dv.add(f"E6:E{5 + NF}"); fl.add_data_validation(dv)
name("FlDate", f"Flights!$A$6:$A${5 + NF}"); name("FlShift", f"Flights!$B$6:$B${5 + NF}"); name("FlCode", f"Flights!$C$6:$C${5 + NF}")
name("FlTime", f"Flights!$D$6:$D${5 + NF}"); name("FlKey", f"Flights!$I$6:$I${5 + NF}"); name("FlOrder", f"Flights!$J$6:$J${5 + NF}"); name("FlSorties", f"Flights!$G$6:$G${5 + NF}"); name("FlCount", f"Flights!$K$6:$K${5 + NF}"); name("FlHrs", f"Flights!$L$6:$L${5 + NF}")
fl.sheet_properties.tabColor = "1D4ED8"

# ---- Day Status: one row per date
banner(days, "DAY STATUS", "Leave blank for a flying day. Pick a reason (No Flying, Stand-down, Maintenance ...) for a day with no flying: it greys out on the calendar and departments see NO FLY.", 3)
header(days, 5, ["Date", "Day", "Status"], {3: "B91C1C"})
status_sample = seedplan.STATUS
d0 = dt.date(YEAR, 1, 1); ndays = 397   # row 6 + day number from 1 Jan of the planning year, through 31 Jan of the next year (the look-ahead)
for i in range(ndays):
    r = 6 + i; dte = d0 + dt.timedelta(days=i)
    a = days.cell(r, 1, f'=IF(DATE(PlanYear,1,1)+{i}<=DATE(PlanYear+1,1,31),DATE(PlanYear,1,1)+{i},"")'); a.number_format = "dd mmm yyyy"; a.font = f(False, SLATE700)
    b = days.cell(r, 2, f'=IF(A{r}="","",TEXT(A{r},"ddd"))'); b.font = f(False, SLATE500); b.alignment = CENTER
    cst = days.cell(r, 3, status_sample.get(dte)); cst.fill = fill("FFFBEB"); cst.font = f(True, "B91C1C")
    for cc in range(1, 4): days.cell(r, cc).border = BOX
days.conditional_formatting.add(f"A6:B{5 + ndays}", FormulaRule(formula=['AND($A6<>"",WEEKDAY($A6,2)>5)'], fill=fill(SLATE100), font=Font(name=FONT, bold=True, color=SLATE700)))
dv = DataValidation(type="list", formula1="DayStatusList", allow_blank=True, showErrorMessage=True, errorStyle="information", errorTitle="Day status", error="Not on the list, but it will still mark the day as no flying.")
dv.add(f"C6:C{5 + ndays}"); days.add_data_validation(dv)
for col, w in zip("ABC", [16, 8, 24]): days.column_dimensions[col].width = w
days.freeze_panes = "A6"
days.conditional_formatting.add(f"A6:C{5 + ndays}", FormulaRule(formula=['$C6<>""'], fill=fill("FEE2E2")))
name("DayStat", f"'Day Status'!$C$6:$C${5 + ndays}")

# ---- Ground Programme: office / ground events per day (any number), shown on the calendar under the date
NG = 500
gp = wb.create_sheet("Ground Programme", wb.sheetnames.index("Day Status") + 1)
banner(gp, "GROUND PROGRAMME", "Ground events for the office: lectures, briefs, sim, meetings. One per row, any number per day. Crew blank = everyone; or tags P L T S1 S2.", 5)
header(gp, 5, ["Date", "Time", "Activity", "Crew (blank = all)", "Comment"], {3: "15803D"})
gsample = seedplan.GROUND
for i in range(NG):
    r = 6 + i
    for c in range(1, 6):
        cell = gp.cell(r, c); cell.border = BOX; cell.fill = fill("FFFBEB"); cell.font = f(c == 3, "1D4ED8")
    gp.cell(r, 1).number_format = "ddd dd mmm yyyy"; gp.cell(r, 2).number_format = "@"; gp.cell(r, 2).alignment = CENTER
    if i < len(gsample):
        for c, v in enumerate(gsample[i], start=1):
            if v != "": gp.cell(r, c, v)
for col, w in zip("ABCDE", [17, 12, 22, 16, 40]): gp.column_dimensions[col].width = w
gp.freeze_panes = "A6"
dv = DataValidation(type="list", formula1="GrList", allow_blank=True, showErrorMessage=True, errorStyle="information", errorTitle="Ground programme", error="Not on the list, but it will still be shown."); dv.add(f"C6:C{5 + NG}"); gp.add_data_validation(dv)
dv = DataValidation(type="date", operator="between", formula1="DATE(PlanYear,1,1)", formula2="DATE(PlanYear+1,1,31)", allow_blank=True, showErrorMessage=True, errorTitle="Date", error="Use a date in the planning year (Dashboard A5) or January of the next year."); dv.add(f"A6:A{5 + NG}"); gp.add_data_validation(dv)
gp.sheet_properties.tabColor = "15803D"
name("GrDate", f"'Ground Programme'!$A$6:$A${5 + NG}"); name("GrAct", f"'Ground Programme'!$C$6:$C${5 + NG}")

# ---- Requests: the Master Planner's answers to the departments' requests (raised on each department's own Requests sheet, matched by Dept + ID)
NRQ = 500
rq = wb.create_sheet("Requests", wb.sheetnames.index("Ground Programme") + 1)
banner(rq, "REQUESTS FROM DEPARTMENTS", "Answered in the dashboard (Master Planning, Requests inbox). Done or Declined with a reply; the department sees it.", 6)
header(rq, 5, ["Dept", "Request ID", "Status", "Reply", "Answered on", "Request"])
for i in range(NRQ):
    r = 6 + i
    for c in range(1, 7):
        cell = rq.cell(r, c); cell.border = BOX; cell.fill = fill("FFFBEB"); cell.font = f(c == 3, "1D4ED8")
    rq.cell(r, 5).number_format = "dd mmm yyyy"; rq.cell(r, 2).number_format = "@"
for c, v in enumerate([] if REAL else ["SUP 2", "S2-SAMPLE2", "Declined", "Range is booked 0800-1000, keep 0830", dt.date(YEAR, 10, 2), "Raised Thu, 1 Oct: Can Friday's AM flight start 0900 instead? Ground school runs late"], start=1): rq.cell(6, c, v)
for col, w in zip("ABCDEF", [12, 16, 11, 44, 14, 50]): rq.column_dimensions[col].width = w
dv = DataValidation(type="list", formula1='"Done,Declined"', allow_blank=True); dv.add(f"C6:C{5 + NRQ}"); rq.add_data_validation(dv)
rq.freeze_panes = "A6"; rq.sheet_properties.tabColor = "D97706"

# ---- month calendars (view only)
for mi in range(13):   # Jan..Dec + next January
    ws = wb.create_sheet(MONTHS13[mi]); MFULL = MONTHS_FULL[mi % 12]
    weeks = 6   # always six weeks: which cells hold a date follows the year cell (Dashboard A5)
    FIRST = f"DATE(PlanYear,{mi + 1},1)"
    last = 1 + 21
    banner(ws, f"MASTER PLANNER  |  {MFULL.upper()}",
           f"View only: built from the Flights and Day Status sheets. Shows up to {LINES} flights per shift; more show as '+n more'. Plan on the Flights sheet or in the planner dashboard.", last)
    ws.cell(1, 1, f'="MASTER PLANNER  |  {MFULL.upper()} "&(PlanYear+{mi // 12})' + (' & "  (NEXT YEAR)"' if mi == 12 else ""))
    if mi == 12: ws.sheet_properties.tabColor = "7C3AED"
    ws.row_dimensions[3].height = 22
    ws.cell(3, 2, "Planned sorties this month").font = f(True, SLATE700, 10)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=4)
    tot = ws.cell(3, 5); tot.font = f(True, "1D4ED8", 14); tot.alignment = CENTER
    rng = f'FlDate,">="&{FIRST},FlDate,"<="&EOMONTH({FIRST},0)'
    for c0, lab, fx, nf in ((7, "Flights", f"=SUMIFS(FlCount,{rng})", "0"), (10, "Planned flying hrs", f"=SUMIFS(FlHrs,{rng})", "0.0")):
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 1)
        ws.cell(3, c0, lab).font = f(True, SLATE700, 10)
        v = ws.cell(3, c0 + 2, fx); v.font = f(True, "0F766E" if c0 == 10 else "1D4ED8", 14); v.alignment = CENTER; v.number_format = nf
    ws.cell(3, 14, "To add, change or remove a flight: use the Flights sheet, or the Master Planning tab in the planner dashboard.").font = f(False, SLATE500, 9, True)
    for k, dn in enumerate(["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]):
        c0 = 2 + 3 * k
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
        for c in range(c0, c0 + 3): ws.cell(4, c).fill = fill(SLATE700 if k > 4 else SLATE900)
        h = ws.cell(4, c0, dn.upper()); h.font = f(True, "FCD34D" if k > 4 else "FFFFFF", 11); h.alignment = CENTER
    ws.row_dimensions[4].height = 22
    day_counts = []
    for w in range(weeks):
        r0 = 5 + WEEK_ROWS * w
        for k in range(7):
            c0 = 2 + 3 * k; n0 = 7 * w + k; inside = True   # every cell gets formulas; outside the month the date is blank and the cell greys out
            dref = f"${CL(c0)}${r0}"
            ws.conditional_formatting.add(f"{CL(c0)}{r0}:{CL(c0 + 2)}{r0 + WEEK_ROWS - 1}", FormulaRule(formula=[f'{dref}=""'], fill=fill(SLATE100), font=Font(color=SLATE100), stopIfTrue=True))
            ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 1)
            ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 1, end_column=c0 + 2)
            for r in (r0, r0 + 1):
                for c in range(c0, c0 + 3):
                    ws.cell(r, c).fill = fill((("E2E8F0" if k > 4 else "DBEAFE") if r == r0 else "FFFFFF") if inside else SLATE100)
                    ws.cell(r, c).border = Border(left=heavy if c == c0 else None, right=heavy if c == c0 + 2 else None, top=heavy if r == r0 else hair, bottom=thin)
            if inside:
                dx = f"{FIRST}-WEEKDAY({FIRST},3)+{n0}"
                h = ws.cell(r0, c0, f'=IF(MONTH({dx})={mi % 12 + 1},{dx},"")'); h.number_format = 'ddd d mmm'; h.font = f(True, SLATE900, 11); h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                sref = f"${CL(c0)}${r0 + 1}"
                stx = f'IFERROR(INDEX(DayStat,{dref}-DATE(PlanYear,1,1)+1)&"","")'   # the day status (blank = flying day)
                cnt = ws.cell(r0, c0 + 2, f'=IF({dref}="","",IF({stx}<>"",0,SUMIFS(FlSorties,FlDate,{dref})))'); cnt.number_format = '[=1]0" sortie";[>1]0" sorties";"-"'
                cnt.font = f(True, "1D4ED8", 9); cnt.alignment = CENTER
                day_counts.append(f"{CL(c0 + 2)}{r0}")
                # under the date: the no-flying reason, else the ground programme ("GND: Safety Brief +1 more")
                gn = f'COUNTIF(GrDate,{dref})'
                st = ws.cell(r0 + 1, c0, f'=IF({dref}="","",IF({stx}<>"",{stx},IF({gn}=0,"","GND: "&INDEX(GrAct,MATCH({dref},GrDate,0))&IF({gn}>1," +"&({gn}-1)&" more",""))))'); st.font = f(True, "166534", 9); st.alignment = CENTER
                ws.conditional_formatting.add(f"{CL(c0)}{r0 + 1}", FormulaRule(formula=[f'{stx}<>""'], fill=fill("B91C1C"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
                ws.conditional_formatting.add(f"{CL(c0)}{r0 + 1}", FormulaRule(formula=[f'{sref}<>""'], fill=fill("DCFCE7"), font=Font(color="166534", bold=True)))
            for s, sh in enumerate(SHIFTS):
                rs = r0 + 2 + LINES * s
                bg, fg = SHIFT_STYLE[s]
                ws.merge_cells(start_row=rs, start_column=c0, end_row=rs + LINES - 1, end_column=c0)
                for i in range(LINES):
                    for c in range(c0, c0 + 3):
                        cell = ws.cell(rs + i, c)
                        cell.fill = fill(bg if c == c0 else "FFFFFF") if inside else fill(SLATE100)
                        bottom = (heavy if s == 2 else Side(style="thin", color=SLATE400)) if i == LINES - 1 else hair
                        cell.border = Border(left=heavy if c == c0 else hair, right=heavy if c == c0 + 2 else hair, top=hair, bottom=bottom)
                if not inside: continue
                n = f'COUNTIFS(FlDate,{dref},FlShift,"{sh}",FlCode,"<>")'
                lab = ws.cell(rs, c0, f'=IF({dref}="","","{sh}"&IF({n}>{LINES},CHAR(10)&"+"&({n}-{LINES})&" more",""))')
                lab.font = Font(name="Consolas", bold=True, color=fg, size=7); lab.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                for i in range(LINES):
                    k_ = f'MATCH({dref}&"|{sh}|{i + 1}",FlKey,0)'
                    cc = ws.cell(rs + i, c0 + 1, f'=IF({dref}="","",IFERROR(INDEX(FlCode,{k_}),""))'); cc.font = f(True, "166534", 9); cc.alignment = CENTER
                    tc = ws.cell(rs + i, c0 + 2, f'=IF({dref}="","",IFERROR(INDEX(FlTime,{k_})&"",""))'); tc.font = f(False, SLATE700, 8); tc.alignment = CENTER
                ws.conditional_formatting.add(f"{CL(c0 + 1)}{rs}:{CL(c0 + 1)}{rs + LINES - 1}", FormulaRule(formula=[f'AND({CL(c0 + 1)}{rs}<>"",{stx}<>"")'], fill=fill("FCA5A5"), font=Font(color="7F1D1D", bold=True, strike=True), stopIfTrue=True))
                ws.conditional_formatting.add(f"{CL(c0 + 1)}{rs}:{CL(c0 + 1)}{rs + LINES - 1}", FormulaRule(formula=[f'{CL(c0 + 1)}{rs}<>""'], fill=fill("DCFCE7")))
            if inside:
                ws.conditional_formatting.add(f"{CL(c0)}{r0 + 2}:{CL(c0 + 2)}{r0 + WEEK_ROWS - 1}", FormulaRule(formula=[f'{stx}<>""'], fill=fill("E5E7EB"), font=Font(color=SLATE400)))
        ws.row_dimensions[r0].height = 19; ws.row_dimensions[r0 + 1].height = 14
        for i in range(3 * LINES): ws.row_dimensions[r0 + 2 + i].height = 14
    tot.value = "=" + "+".join(f"N({x})" for x in day_counts)
    ws.column_dimensions["A"].width = 2
    for k in range(7):
        c0 = 2 + 3 * k
        ws.column_dimensions[CL(c0)].width = 7; ws.column_dimensions[CL(c0 + 1)].width = 10.5; ws.column_dimensions[CL(c0 + 2)].width = 10
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 85; ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "4:4"
    ws.sheet_properties.tabColor = "64748B"

# ---- Dashboard: year + file name (the planner dashboard reads A5 and G5) and sorties per month
banner(dash, "MASTER PLANNER", "The master list of planned flights for every department.", 9)
dash["A4"] = "YEAR"; dash["A4"].font = f(True, SLATE500, 9)
dash["A5"] = YEAR; dash["A5"].font = f(True, "1D4ED8", 14)
dash["G4"] = "FILE"; dash["G4"].font = f(True, SLATE500, 9)
dash["G5"] = "Master Planner"; dash["G5"].font = f(True, "BE123C", 14)
dash["A7"] = "Month"; dash["B7"] = "Planned sorties"; dash["C7"] = "Flights"; dash["D7"] = "Planned hrs"
for c in ("A7", "B7", "C7", "D7"): dash[c].font = f(True, "FFFFFF"); dash[c].fill = fill(SLATE900); dash[c].alignment = CENTER
for mi in range(12):
    a = dash.cell(8 + mi, 1, MONTHS_FULL[mi]); a.hyperlink = f"#'{MONTHS[mi]}'!A1"; a.font = f(color="1D4ED8")
    b = dash.cell(8 + mi, 2, f"='{MONTHS[mi]}'!E3"); b.alignment = CENTER; b.font = f(True)
    b = dash.cell(8 + mi, 3, f"='{MONTHS[mi]}'!I3"); b.alignment = CENTER; b.font = f(True)
    b = dash.cell(8 + mi, 4, f"='{MONTHS[mi]}'!L3"); b.alignment = CENTER; b.font = f(True); b.number_format = "0.0"
dash.cell(20, 1, "Total").font = f(True)
for col in "BCD": dash[f"{col}20"] = f"=SUM({col}8:{col}19)"; dash[f"{col}20"].font = f(True); dash[f"{col}20"].alignment = CENTER
dash["D20"].number_format = "0.0"
a = dash.cell(21, 1, "Jan (next year)"); a.hyperlink = "#'Jan (next year)'!A1"; a.font = f(color="7C3AED")
for k, col in enumerate("EIL"):
    b = dash.cell(21, 2 + k, f"='Jan (next year)'!{col}3"); b.alignment = CENTER; b.font = f(True, "7C3AED")
dash["D21"].number_format = "0.0"
tips = ["HOW IT WORKS",
        "1. Flights sheet: the master list, one flight per row (date, AM / PM / NIGHT, flight code, time, comment). Any number of flights per shift.",
        "2. Easiest is the Master Planning tab in the planner dashboard: + Add flight under a shift, x to remove one.",
        "3. Crew comes from the flight code (Flight Codes sheet, column E: P Pilot, L Loadmaster, T TACCO, S1 SUP 1, S2 SUP 2; blank = everyone).",
        "   Type a crew on a flight only when it differs from its code. Departments only get the flights that need them.",
        "4. No flying on a day: pick a reason on the Day Status sheet (or in the dashboard). The day turns grey and departments see NO FLY.",
        "5. Jan to Dec tabs (and Jan (next year), for planning January while still in December): a calendar view of the lists, with each day's sorties. View only.",
        "6. Departments take their flights from here: in the planner dashboard, open your department file with the Master Planner and click 'Copy flights from the Master Planner'."]
for k, t in enumerate(tips):
    c = dash.cell(8 + k, 6, t); c.font = f(k == 0, SLATE900 if k == 0 else SLATE700, 11 if k == 0 else 10)
dash.column_dimensions["A"].width = 14; dash.column_dimensions["B"].width = 16; dash.column_dimensions["C"].width = 10; dash.column_dimensions["D"].width = 13; dash.column_dimensions["E"].width = 3
wb.active = wb.sheetnames.index("Oct")
# layout stamp: the dashboard compares it with the layout it was written for (bump it when cell positions change)
dash["Z2"] = "LODESTAR master layout 2"; dash.column_dimensions["Z"].hidden = True
wb.save(OUT)
print("saved", OUT)
