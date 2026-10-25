import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

# Ops Room file: kept by the Ops Room duty staff (SDO and AOS). The flights come from the Master Planner and the crew from the
# department planners; this file only holds what the Ops Room decides per flight: seat changes (blank = the dashboard's pick),
# the Mission Commander and the CRMO. One row per flight, found by date + shift + code + time + # (the nth identical flight that day).
OUT = sys.argv[1]
YEAR = 2026
FONT = "Arial"
N = 2000
SEATS = ["Aircraft Captain", "Co-Pilot", "Loadmaster", "TACCO", "SUP 1", "SUP 2"]
SLATE900, SLATE700, SLATE500, SLATE400 = "0F172A", "334155", "64748B", "94A3B8"
def f(bold=False, color="0F172A", size=10, italic=False): return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="CBD5E1"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True); LEFT = Alignment(horizontal="left", vertical="center", indent=1)
INPUT_FILL = fill("FFFBEB")
wb = Workbook()
def banner(ws, title, subtitle, last_col):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        for c in range(1, last_col + 1): ws.cell(r, c).fill = fill(SLATE900)
    ws.cell(1, 1, title).font = f(True, "FFFFFF", 16); ws.cell(2, 1, subtitle).font = f(False, SLATE400, 10)
    ws.cell(1, 1).alignment = LEFT; ws.cell(2, 1).alignment = LEFT
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 18

dash = wb.active; dash.title = "Dashboard"
banner(dash, "OPS ROOM  |  FLIGHTS OF THE DAY", "Kept by the Ops Room duty staff (SDO and AOS). Flights come from the Master Planner, crew from each department's planner.", 8)
dash["A4"] = "YEAR"; dash["A4"].font = f(True, SLATE500, 9)
dash["A5"] = YEAR; dash["A5"].font = f(True, "1D4ED8", 14); dash["A5"].fill = INPUT_FILL; dash["A5"].border = BOX; dash["A5"].alignment = CENTER
dv = DataValidation(type="list", formula1='"' + ",".join(str(y) for y in range(2025, 2036)) + '"'); dv.add("A5"); dash.add_data_validation(dv)
dash["G4"] = "FILE"; dash["G4"].font = f(True, SLATE500, 9)
dash["G5"] = "Ops Room"; dash["G5"].font = f(True, "0F766E", 14)
wb.defined_names["PlanYear"] = DefinedName("PlanYear", attr_text="Dashboard!$A$5")
steps = [("HOW THE OPS ROOM WORKS", True),
         ("1. Open the planner dashboard and click Ops Room. It opens every department's planner, the Master Planner, the Duty Roster and this file.", False),
         ("2. The Ops Room tab shows the day's flights from the Master Planner, the ground programme and tomorrow's forecast.", False),
         ("3. Each flight lists its crew from the department plans: Aircraft Captain, Co-Pilot (Pilot), Loadmaster, TACCO, SUP 1 and SUP 2. Of two pilots, the more senior is suggested as Captain; swap them if needed.", False),
         ("4. Mark the Mission Commander (MC) and the CRMO on each flight. Save to Excel writes them to this file.", False),
         ("", False),
         ("Flight Crew sheet: one row per aircraft the Ops Room has changed (callsign, tail, AO, times, crew, MC, CRMO, remarks). A blank seat means the dashboard's own pick from the department plans.", False)]
for i, (t, h) in enumerate(steps):
    dash.merge_cells(start_row=8 + i, start_column=1, end_row=8 + i, end_column=8)
    c = dash.cell(8 + i, 1, t); c.font = f(True, SLATE900, 12) if h else f(False, SLATE700, 11)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1); dash.row_dimensions[8 + i].height = 22 if h or not t else 32
for col, w in zip("ABCDEFGH", [12, 14, 14, 14, 14, 14, 16, 14]): dash.column_dimensions[col].width = w

fc = wb.create_sheet("Flight Crew")
# one row per aircraft: a 2x code (formation) has an aircraft 1 and an aircraft 2 row
heads = ["Date", "Shift", "Flight code", "Time", "#", "Aircraft", "Callsign", "Tail"] + SEATS + ["Others", "AO", "ETD", "ETA", "ATD", "ATA", "Mission Commander", "CRMO", "Remarks", "Alert Status", "Status", "Cancel Reason", "Standby Tail", "Area Time", "Additional Crew", "Pax", "Brief Timing"]   # Brief Timing (AE): set by the Ops Room   # Status: Cancelled (Flown = ATD and ATA filled)
banner(fc, "FLIGHT CREW  |  OPS ROOM DECISIONS", "Filled in by the dashboard's Ops Room tab, one row per aircraft. Crew are callsigns; a blank seat means the dashboard's pick from the department plans. Blank ETD / ETA = the Master Planner time.", len(heads))
for i, h in enumerate(heads):
    c = fc.cell(5, 1 + i, h); c.font = f(True, "FFFFFF", 10); c.fill = fill(SLATE900); c.alignment = CENTER; c.border = BOX
widths = [12, 7, 12, 11, 4, 8, 11, 7, 14, 12, 12, 10, 10, 10, 14, 9, 7, 7, 7, 7, 14, 10, 30, 12, 10, 16, 10, 24, 30, 16, 10]
from openpyxl.utils import get_column_letter as CL
for i, w in enumerate(widths): fc.column_dimensions[CL(1 + i)].width = w
for r in range(6, 6 + N):
    fc.cell(r, 1).number_format = "dd mmm yyyy"
    fc.cell(r, 28).alignment = Alignment(wrap_text=True, vertical="top"); fc.cell(r, 23).alignment = Alignment(wrap_text=True, vertical="top")
import realdata
REAL = realdata.ON   # real squadron data: no sample Ops Room rows
TAILS = realdata.TAILS if REAL and realdata.TAILS else [str(t) for t in range(710, 719)]
for k, (cs, tail) in enumerate(() if REAL else (("RICOH 1", "716"), ("RICOH 2", "712"))):
    r = 6 + k
    for c, v in enumerate([dt.date(YEAR, 10, 1), "AM", "2xFLT-2.0", "0830-1030", 1, k + 1, cs, tail], start=1): fc.cell(r, c, v)
    fc.cell(r, 16, "Area 1"); fc.cell(r, 31, "0730"); fc.cell(r, 27, ("714", "715")[k])
if not REAL: fc.cell(6, 28, "Oscar 0840-0950\nEcho 0950-1010"); fc.cell(6, 29, "SUP 2: OWL"); fc.cell(6, 30, "OWL"); fc.cell(6, 21, "COBRA"); fc.cell(6, 22, "VIPER"); fc.cell(6, 23, "Formation brief 0730")
sv = DataValidation(type="list", formula1='"AM,PM,NIGHT"', allow_blank=True); sv.add(f"B6:B{5 + N}"); fc.add_data_validation(sv)
# the squadron's aircraft (tails 710-718): actual tail (H) and standby tail (AA), picked in the Ops Room
tv = DataValidation(type="list", formula1='"' + ",".join(TAILS) + '"', allow_blank=True); tv.add(f"H6:H{5 + N}"); tv.add(f"AA6:AA{5 + N}"); fc.add_data_validation(tv)
fc.freeze_panes = "A6"; fc.sheet_properties.tabColor = "0F766E"
# Day Board: who holds each duty that day (the header above the flight table). Row = 6 + day of the year.
DB_HEADS = ["Date", "SXO", "OPS O", "SDO", "BDO", "Safety", "Prefect", "Security", "Intel 1", "Intel 2", "DOB", "Oscar", "Echo", "Nov", "CODES", "Remarks", "Standby Sets"]   # Standby Sets (Q): 1 = one crew covers P1 and P2, 2 or blank = separate crews   # Remarks (P): the Ops Room remarks card under the ground programme
db = wb.create_sheet("Day Board")
banner(db, "DAY BOARD  |  DUTY PERSONNEL", "Filled in by the dashboard's Ops Room tab: one row per day. Blank SXO, OPS O, SDO and BDO = whoever has that ground duty planned in the departments' plans that day.", len(DB_HEADS))
for i, h in enumerate(DB_HEADS):
    c = db.cell(5, 1 + i, h); c.font = f(True, "0F172A", 10); c.fill = fill("FACC15"); c.alignment = CENTER; c.border = BOX
d0 = dt.date(YEAR, 1, 1); nd = 397   # row 6 + day number from 1 Jan of the planning year (Dashboard A5), through 31 Jan of the next year (the look-ahead)
for k in range(nd):
    c = db.cell(6 + k, 1, f'=IF(DATE(PlanYear,1,1)+{k}<=DATE(PlanYear+1,1,31),DATE(PlanYear,1,1)+{k},"")'); c.number_format = "ddd dd mmm yyyy"; c.font = f(True, SLATE700, 9)
    for j in range(2, len(DB_HEADS) + 1): db.cell(6 + k, j).border = BOX
r = 6 + (dt.date(YEAR, 10, 1) - d0).days
for j, v in enumerate([None, None, None, None, "CPT Lim", "ME2 Goh", "ME1 Ravi", "CPT Lee", "LTA Ng", "0700", "Oscar 1", "Echo 2", "Nov 3", "Alpha / Kilo", "Wx brief 0700. Range Area 1 NOTAM active 0800-1200."], start=2):
    if v and not REAL: db.cell(r, j, v)
db.column_dimensions["A"].width = 18
for j in range(2, len(DB_HEADS) + 1): db.column_dimensions[CL(j)].width = 12
db.column_dimensions[CL(len(DB_HEADS) - 1)].width = 50; db.column_dimensions[CL(len(DB_HEADS))].width = 13
db.freeze_panes = "B6"; db.sheet_properties.tabColor = "FACC15"
# Medical Chits: up-chits the Ops Room records for a department whose own file has no Medical Chits sheet (older layout). Rows 6..505.
mc = wb.create_sheet("Medical Chits")
MC_HEADS = ["Department", "Callsign", "Down-chit", "Up-chit", "Up-chitted by", "Remarks"]
banner(mc, "MEDICAL CHITS  |  RECORDED IN THE OPS ROOM", "Up-chits made in the Ops Room for departments whose own planner has no Medical Chits sheet yet. Up-chit date = first day fit to fly.", len(MC_HEADS))
for i, h in enumerate(MC_HEADS):
    c = mc.cell(5, 1 + i, h); c.font = f(True, "FFFFFF", 10); c.fill = fill(SLATE900); c.alignment = CENTER; c.border = BOX
for r in range(6, 506):
    mc.cell(r, 3).number_format = "ddd dd mmm yyyy"; mc.cell(r, 4).number_format = "ddd dd mmm yyyy"
for col, w in zip("ABCDEF", [14, 12, 16, 16, 14, 40]): mc.column_dimensions[col].width = w
mc.freeze_panes = "A6"; mc.sheet_properties.tabColor = "DC2626"
# Ops Add: ad hoc flights the Ops Room adds on the day. Not on the Master Planner; departments see them on their MASTER rows (comment "OPS ADD"). Rows 6..505.
oa = wb.create_sheet("Ops Add", 1)
OA_HEADS = ["Date", "Shift", "Flight code", "Time", "Crew", "Comment"]
banner(oa, "OPS ADD  |  AD HOC FLIGHTS", "Flights added by the Ops Room on the day. They are not on the Master Planner; the dashboard adds them to the departments' programmes. Crew: P Pilot, L Loadmaster, T TACCO, S1 SUP 1, S2 SUP 2 (blank = the flight code's crew).", len(OA_HEADS))
for i, h in enumerate(OA_HEADS):
    c = oa.cell(5, 1 + i, h); c.font = f(True, "FFFFFF", 10); c.fill = fill("7C3AED"); c.alignment = CENTER; c.border = BOX
for r in range(6, 506): oa.cell(r, 1).number_format = "ddd dd mmm yyyy"
for col, w in zip("ABCDEF", [16, 8, 13, 12, 12, 40]): oa.column_dimensions[col].width = w
sv = DataValidation(type="list", formula1='"AM,PM,NIGHT"', allow_blank=True); sv.add("B6:B505"); oa.add_data_validation(sv)
oa.freeze_panes = "A6"; oa.sheet_properties.tabColor = "7C3AED"
# layout stamp: the dashboard compares it with the layout it was written for (bump it when cell positions change)
dash["Z2"] = "LODESTAR ops layout 2"; dash.column_dimensions["Z"].hidden = True
wb.save(OUT); print("saved", OUT, wb.sheetnames)
