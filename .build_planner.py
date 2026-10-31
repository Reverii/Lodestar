import os, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import CustomFilter, CustomFilters, FilterColumn
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

OUT = sys.argv[1]
DEPT = sys.argv[2] if len(sys.argv) > 2 else "SUP 2"   # the department that owns this file
NOFILTER = os.environ.get("NOFILTER") == "1"
NOCHART = os.environ.get("NOCHART") == "1"; NOCOMMENT = os.environ.get("NOCOMMENT") == "1"; STATICNAMES = os.environ.get("STATICNAMES") == "1"; SMALLFREEZE = os.environ.get("SMALLFREEZE") == "1"; WEBCLEAN = os.environ.get("WEBCLEAN") == "1"; COMPACT = os.environ.get("COMPACT") == "1"   # Excel Online test variants
YEAR = 2026
N = 60                # crew slots
NA = 60               # activity slots
HAS_DUTY = True   # BDO (every department) + SDO (SDO_DEPTS) duty days from the Duty Roster
HAS_GND = True   # ground programme row (from the Master Planner)
GND_ROW = 24 if HAS_GND else None
G0 = 27 if HAS_GND else 26   # first grid row on month tabs (6-23 master slots, [24 ground], SDO dept, filter header)
SDO_DEPT_ROW = G0 - 2
SDO_DEPTS = ["SUP 1", "SUP 2", "Loadmaster"]   # only these departments are rostered for SDO
DOES_SDO = DEPT in SDO_DEPTS
FILTER_ROW = G0 - 1
CODE_ROW = lambda s, k: 6 + 6 * s + 3 * k   # master code row for shift s, slot k; time row is +1, comment row is +2
PER = 4              # rows per person on the month tabs: AM, PM, NIGHT and WD (whole day, ground and admin items only)
GL = G0 + PER * N - 1   # last grid row
D1, D31 = 4, 34       # day columns D..AH
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MON_NX = "Jan (next year)"   # look-ahead: January of the next year, planned while still in December (Start new year moves it into Jan)
MONTHS13 = MONTHS + [MON_NX]
# a date's month tab number 1..13 (13 = January of the next year) and "the date is in the plan year or next January"
MIX = lambda a: f"IF(YEAR({a})>PlanYear,13,MONTH({a}))"
INYR = lambda a: f"AND({a}>=DATE(PlanYear,1,1),{a}<=DATE(PlanYear+1,1,31))"
MONTHS_FULL = ["January", "February", "March", "April", "May", "June", "July", "August",
               "September", "October", "November", "December"]
FONT = "Arial"

SLATE900, SLATE800, SLATE700, SLATE500, SLATE400 = "0F172A", "1E293B", "334155", "64748B", "94A3B8"
SLATE200, SLATE100, SLATE50 = "E2E8F0", "F1F5F9", "F8FAFC"
AMBER300 = "FCD34D"

def f(bold=False, color="0F172A", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="CBD5E1")
heavy = Side(style="medium", color=SLATE400)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
INPUT_FILL, INPUT_FONT = fill("FFFBEB"), f(color="1D4ED8")

CATS = [
    ("Flight", "DBEAFE", "1E3A8A", False),
    ("Sim", "CFFAFE", "164E63", False),
    ("Ground Duty", "FEF3C7", "78350F", False),
    ("Squadron Duty", "FECDD3", "881337", False),
    ("Leave", SLATE200, SLATE500, True),
    ("Off", SLATE100, SLATE400, False),
    ("Admin", "EDE9FE", "5B21B6", False),
    ("Medical", "FCE7F3", "9D174D", False),
    ("Oversea Duty", "DCFCE7", "14532D", False),
]
CAT_ROW = {c[0]: 6 + i for i, c in enumerate(CATS)}
METRICS = [
    ("Flights", "Flight", "entries", "172554"),
    ("Fly Hrs", "Flight", "hours", "042F2E"),
    ("Gnd Duty", "Ground Duty", "days", "451A03"),
    ("Sqn Duty", "Squadron Duty", "days", "4C0519"),
    ("Leave", "Leave", "days", SLATE800),
]
MC0 = D31 + 1
DUTY_CODES = ("SDO", "BDO") if HAS_DUTY else ("SDO",)
LASTC = MC0 + len(METRICS) - 1

wb = Workbook()
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
def catref(cat): return f"Activities!$G${CAT_ROW[cat]}"

def banner(ws, title, subtitle, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    for r in (1, 2):
        for c in range(1, last_col + 1):
            ws.cell(r, c).fill = fill(SLATE900)
    ws.cell(1, 1, title).font = f(True, "FFFFFF", 16)
    ws.cell(2, 1, subtitle).font = f(False, SLATE400, 10)
    ws.cell(1, 1).alignment = LEFT; ws.cell(2, 1).alignment = LEFT
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 18

def header(ws, row, col, labels, color=SLATE900):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col + i, lab)
        c.font = f(True, "FFFFFF", 10); c.fill = fill(color); c.alignment = CENTER; c.border = BOX

def card(ws, row, c0, width, label, formula, accent, numfmt="0"):
    ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c0 + width - 1)
    ws.merge_cells(start_row=row + 1, start_column=c0, end_row=row + 1, end_column=c0 + width - 1)
    l = ws.cell(row, c0, label.upper()); l.font = f(True, SLATE500, 9); l.alignment = LEFT
    v = ws.cell(row + 1, c0, formula); v.font = f(True, accent, 18); v.alignment = LEFT; v.number_format = numfmt
    for c in range(c0, c0 + width):
        ws.cell(row, c).fill = fill("FFFFFF"); ws.cell(row + 1, c).fill = fill("FFFFFF")
        ws.cell(row, c).border = Border(top=Side(style="thick", color=accent), left=thin if c == c0 else None, right=thin if c == c0 + width - 1 else None)
        ws.cell(row + 1, c).border = Border(bottom=thin, left=thin if c == c0 else None, right=thin if c == c0 + width - 1 else None)
    ws.row_dimensions[row + 1].height = 30

def apply_cat_cf(ws, rng, tl):
    for cat, bg, fg, it in CATS:
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND({tl}<>"",IFERROR(INDEX(ActCat,MATCH({tl},ActCodesAll,0)),"")={catref(cat)})'],
            fill=fill(bg), font=Font(color=fg, bold=not it, italic=it), stopIfTrue=True))

def nonblank_filter(ws, ref, col_id):
    """Auto filter that hides empty crew slots. Data > Reapply shows newly added people."""
    if NOFILTER:
        return
    ws.auto_filter.ref = ref
    ws.auto_filter.filterColumn.append(FilterColumn(colId=col_id, customFilters=CustomFilters(customFilter=[CustomFilter(operator="notEqual", val=" ")])))

guide = wb.active; guide.title = "Guide"
dash = wb.create_sheet("Dashboard")
daily = wb.create_sheet("Daily View")
sdo = wb.create_sheet("SDO Plan")
bdo = wb.create_sheet("BDO Plan") if HAS_DUTY else None
roster = wb.create_sheet("Roster")
acts = wb.create_sheet("Activities")

# ---------------------------------------------------------------- Roster
REG_N = 45            # roster rows 6-50 are regular crew, rows 51-65 are SANS / NSmen (shown at the bottom of every planner)
SANS_ROW0 = 6 + REG_N
CREW_CATS = [("CAT A", "000000"), ("CAT B", "15803D"), ("CAT C", "D4A800"), ("CAT D", "DC2626"), ("TRAINEE", "C026D3")]
ROLES = ["Instructor", "Senior CAT A", "Junior"]
SANS_FILL = "F3E8FF"
banner(roster, f"CREW ROSTER  |  {DEPT.upper()}", f"Rows 6-{5 + REG_N}: regular crew. Rows {SANS_ROW0}-{5 + N} (purple): SANS / NSmen, shown at the bottom of the planner. To remove someone set Status to Inactive. Do not delete or re-order rows: the month tabs link to this list by row.", 7)
header(roster, 5, 1, ["#", "Callsign", "Crew Member (with rank)", "Initials", "Cat Status", "Role", "Status", "Remarks"])
# (callsign, name, cat, role, status); SANS / NSmen go into the bottom section
import seedplan, realdata
SAMPLE_CREW = seedplan.SAMPLE_CREW
REAL = realdata.ON   # real squadron data from LODESTAR_Data_Input.xlsx: no sample plans, notes, requests or syllabus results
# slot index -> crew tuple: regulars from the top, SANS / NSmen from row 51
crew = {}
_reg = [c for c in SAMPLE_CREW.get(DEPT, []) if c[4] not in ("SANS", "NSMEN")]
_ns = [c for c in SAMPLE_CREW.get(DEPT, []) if c[4] in ("SANS", "NSMEN")]
for k, c in enumerate(_reg): crew[k] = c
for k, c in enumerate(_ns): crew[REG_N + k] = c
_real = {}   # slot -> the person's row from the data input (initials, remarks)
if REAL:
    _rp = realdata.CREW.get(DEPT, [])
    for k, p in enumerate([p for p in _rp if p['status'] not in ('SANS', 'NSMEN')]): _real[k] = p
    for k, p in enumerate([p for p in _rp if p['status'] in ('SANS', 'NSMEN')]): _real[REG_N + k] = p
is_sans = lambda i: i >= REG_N
for i in range(N):
    r = 6 + i
    roster.cell(r, 1, i + 1).font = f(color=SLATE400); roster.cell(r, 1).alignment = CENTER
    if is_sans(i): roster.cell(r, 1).fill = fill(SANS_FILL)
    for c in range(2, 9):
        cell = roster.cell(r, c); cell.fill = fill(SANS_FILL) if is_sans(i) else INPUT_FILL; cell.font = INPUT_FONT; cell.border = BOX
    if i in crew:
        cs_, nm_ = crew[i][:2]
        ini_ = "".join(w[0] for w in nm_.replace("(NS)", "").split()[1:] if w[0].isalpha()).upper()   # sample initials from the name, skipping the rank
        if REAL: ini_ = _real[i]['ini'] or None
        for j, v in enumerate((cs_, nm_, ini_) + tuple(crew[i][2:])):
            roster.cell(r, 2 + j, v if v != "" else None)
        if REAL and _real[i]['rem']: roster.cell(r, 8, _real[i]['rem'])
for c in range(1, 9):
    roster.cell(SANS_ROW0, c).border = Border(left=thin, right=thin, top=Side(style="thick", color="7E22CE"), bottom=thin)
roster.cell(SANS_ROW0 - 1, 10, "Regular crew above").font = f(False, SLATE500, 10, italic=True)
roster.cell(SANS_ROW0, 10, "SANS / NSmen from here").font = f(True, "7E22CE", 9)
if crew and not REAL:
    roster.cell(6, 8, "Sample crew: replace with your own")
for col, w in zip("ABCDEFGHIJ", [5, 12, 28, 9, 11, 14, 10, 30, 2, 22]):
    roster.column_dimensions[col].width = w
roster.freeze_panes = "A6"
dv = DataValidation(type="list", formula1='"' + ",".join(c for c, _ in CREW_CATS) + '"', allow_blank=True); dv.add(f"E6:E{5+N}"); roster.add_data_validation(dv)
dv = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True); dv.add(f"F6:F{5+N}"); roster.add_data_validation(dv)
dv = DataValidation(type="list", formula1='"Active,Inactive"', allow_blank=True, showErrorMessage=True, errorTitle="Regular crew",
                    error=f"SANS and NSmen go in the purple rows {SANS_ROW0}-{5 + N}."); dv.add(f"G6:G{5+REG_N}"); roster.add_data_validation(dv)
dv = DataValidation(type="list", formula1='"SANS,NSMEN,Inactive"', allow_blank=True, showErrorMessage=True, errorTitle="SANS / NSmen rows",
                    error=f"These rows are for SANS and NSmen. Regular (Active) crew go in rows 6-{5 + REG_N}."); dv.add(f"G{SANS_ROW0}:G{5+N}"); roster.add_data_validation(dv)
dv = DataValidation(type="custom", formula1=f"COUNTIF($B$6:$B${5+N},B6)=1", allow_blank=True, showErrorMessage=True,
                    errorTitle="Duplicate callsign", error="That callsign is already used."); dv.add(f"B6:B{5+N}"); roster.add_data_validation(dv)
roster.conditional_formatting.add(f"B6:H{5+N}", FormulaRule(formula=['$G6="Inactive"'], fill=fill(SLATE100), font=Font(color=SLATE400, italic=True, strike=True), stopIfTrue=True))
name("CrewName", f"Roster!$C$6:$C${5+N}")
name("CrewCat", f"Roster!$E$6:$E${5+N}")
name("CrewStatus", f"Roster!$G$6:$G${5+N}")
def cat_cf(ws, rng, idx):
    """crew name font colour from Cat Status; idx is the formula giving the roster position (1-based) of the cell's row"""
    for cat, col in CREW_CATS:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'INDEX(CrewCat,{idx})="{cat}"'], font=Font(color=col, bold=True)))
cat_cf(roster, f"C6:C{5+N}", "ROW()-5")

# ---------------------------------------------------------------- Cell Notes: comments on crew cells (date, shift, callsign), written by the planner dashboard or typed here
NOTE_N = 1500
NOTE_SIDE = Side(style="thin", color="7E22CE")
notes = wb.create_sheet("Cell Notes", wb.sheetnames.index("Roster") + 1)
banner(notes, f"CELL NOTES  |  {DEPT.upper()}", "Comments on crew cells. Add them in the planner dashboard (click a cell, type in Comment), or type a row here. Cells with a comment get a purple outline on the month tabs.", 7)
notes["A3"] = "To remove a comment, clear its row. Callsign can also be the crew member's name if they have no callsign."; notes["A3"].font = f(False, SLATE500, 9, italic=True)
header(notes, 5, 1, ["Date", "Shift", "Callsign", "Crew Member", "Comment", "Activity", "Key"])
for col, w in zip("ABCDEFG", [16, 9, 12, 26, 50, 12, 14]): notes.column_dimensions[col].width = w
_crew_idx = lambda r: f'IFERROR(MATCH(C{r},Roster!$B$6:$B${5+N},0),MATCH(C{r},Roster!$C$6:$C${5+N},0))'
_grid_row = lambda r: f'{G0}+{PER}*({_crew_idx(r)}-1)+MATCH(B{r},{{"AM","PM","NIGHT","WD"}},0)-1'
for i in range(NOTE_N):
    r = 6 + i
    for c in range(1, 8):
        cell = notes.cell(r, c); cell.border = BOX
        cell.fill = INPUT_FILL if c in (1, 2, 3, 5) else fill(SLATE100)
        if c in (1, 2, 3, 5): cell.font = INPUT_FONT
    notes.cell(r, 1).number_format = "ddd dd mmm yyyy"; notes.cell(r, 2).alignment = CENTER; notes.cell(r, 3).number_format = "@"
    notes.cell(r, 5).alignment = Alignment(vertical="center", wrap_text=True)
    notes.cell(r, 4, f'=IF(C{r}="","",IFERROR(INDEX(Roster!$C$6:$C${5+N},{_crew_idx(r)}),"not on the Roster"))')
    notes.cell(r, 6, f'=IF(OR(A{r}="",C{r}="",B{r}=""),"",IFERROR(INDEX(CHOOSE(' + MIX(f"A{r}") + ',' + ",".join(f"'{mon}'!$D${G0}:$AH${GL}" for mon in MONTHS13) + f'),{_grid_row(r)}-{G0}+1,DAY(A{r}))&"",""))').alignment = CENTER
    notes.cell(r, 7, f'=IF(OR(A{r}="",C{r}="",B{r}="",E{r}=""),"",IFERROR(A{r}&"|"&({_grid_row(r)}),""))').font = f(False, SLATE400, 8)
notes.freeze_panes = "A6"
dv = DataValidation(type="list", formula1='"AM,PM,NIGHT,WD"', allow_blank=True); dv.add(f"B6:B{5+NOTE_N}"); notes.add_data_validation(dv)
dv = DataValidation(type="date", operator="between", formula1="DATE(PlanYear,1,1)", formula2="DATE(PlanYear+1,1,31)", allow_blank=True, showErrorMessage=True, errorTitle="Date", error=f"Use a date in {YEAR} or January {YEAR + 1}.")
dv.add(f"A6:A{5+NOTE_N}"); notes.add_data_validation(dv)
if not NOFILTER: notes.auto_filter.ref = f"A5:G{5+NOTE_N}"
notes.sheet_properties.tabColor = "7E22CE"
name("NoteKey", f"'Cell Notes'!$G$6:$G${5+NOTE_N}")
if DEPT == "SUP 2" and not REAL:
    for dd, sh, cs, tx in [(1, "AM", "VIPER", "Instructor for EAGLE"), (2, "NIGHT", "VIPER", "NVG currency check"), (5, "AM", "HAWK", "Lead, brief 0730")]:
        r = 6 + [(1, "AM", "VIPER"), (2, "NIGHT", "VIPER"), (5, "AM", "HAWK")].index((dd, sh, cs))
        notes.cell(r, 1, dt.date(YEAR, 10, dd)); notes.cell(r, 2, sh); notes.cell(r, 3, cs); notes.cell(r, 5, tx)

# ---------------------------------------------------------------- Medical Chits: an MC on the plan down-chits the person from that day until an up-chit date here.
# Rows are written by the dashboard (Planner or Ops Room "Up-chit"); a row with a Down-chit date and no Up-chit date down-chits someone without an MC cell.
CHIT_N = 200
chit = wb.create_sheet("Medical Chits", wb.sheetnames.index("Cell Notes") + 1)
banner(chit, f"MEDICAL CHITS  |  {DEPT.upper()}", "MC on the plan = down-chitted from that day (not cleared to fly) until an Up-chit date on this sheet. Up-chit in the planner dashboard (Planner or Ops Room), or type a row here.", 6)
chit["A3"] = "Up-chit date = first day the person is fit to fly again. A row with only a Down-chit date down-chits someone without an MC on the plan."; chit["A3"].font = f(False, SLATE500, 9, italic=True)
header(chit, 5, 1, ["Callsign", "Down-chit", "Up-chit", "Up-chitted by", "Remarks", "Crew Member"])
for col, w in zip("ABCDEF", [12, 16, 16, 14, 40, 26]): chit.column_dimensions[col].width = w
for i in range(CHIT_N):
    r = 6 + i
    for c in range(1, 7):
        cell = chit.cell(r, c); cell.border = BOX
        cell.fill = INPUT_FILL if c <= 5 else fill(SLATE100)
        if c <= 5: cell.font = INPUT_FONT
    chit.cell(r, 2).number_format = "ddd dd mmm yyyy"; chit.cell(r, 3).number_format = "ddd dd mmm yyyy"; chit.cell(r, 1).number_format = "@"
    chit.cell(r, 6, f'=IF(A{r}="","",IFERROR(INDEX(Roster!$C$6:$C${5+N},MATCH(A{r},Roster!$B$6:$B${5+N},0)),"not on the Roster"))')
chit.freeze_panes = "A6"
dv = DataValidation(type="date", operator="between", formula1="DATE(PlanYear,1,1)", formula2="DATE(PlanYear+1,12,31)", allow_blank=True, showErrorMessage=True, errorTitle="Date", error="Use a date.")
dv.add(f"B6:C{5+CHIT_N}"); chit.add_data_validation(dv)
chit.conditional_formatting.add(f"A6:E{5+CHIT_N}", FormulaRule(formula=[f'AND($A6<>"",$B6<>"",$C6="")'], fill=fill("FEE2E2"), font=Font(color="991B1B", bold=True)))
chit.sheet_properties.tabColor = "DC2626"

# ---------------------------------------------------------------- Requests: notes to the Master Planner (add / change / cancel a flight, anything else).
# Written by the dashboard (Master Planning > a day > + Request); the Master Planner answers Done / Declined in its own file, keyed by the ID here.
REQ_N = 200
req = wb.create_sheet("Requests", wb.sheetnames.index("Medical Chits") + 1)
banner(req, f"REQUESTS TO THE MASTER PLANNER  |  {DEPT.upper()}", "Raise them in the dashboard (Master Planning, + New request). The Master Planner's answer shows there.", 7)
req["A3"] = "ID is set by the dashboard and links the answer in the Master Planner file. Clear a row to withdraw a request."; req["A3"].font = f(False, SLATE500, 9, italic=True)
header(req, 5, 1, ["ID", "Date (optional)", "Shift (optional)", "Type (optional)", "Flight code (optional)", "Request", "Raised on"])
for col, w in zip("ABCDEFG", [16, 16, 8, 14, 13, 50, 14]): req.column_dimensions[col].width = w
for i in range(REQ_N):
    r = 6 + i
    for c in range(1, 8):
        cell = req.cell(r, c); cell.border = BOX; cell.fill = INPUT_FILL; cell.font = INPUT_FONT
    req.cell(r, 2).number_format = "ddd dd mmm yyyy"; req.cell(r, 7).number_format = "dd mmm yyyy"; req.cell(r, 1).number_format = "@"
if DEPT == "SUP 2" and not REAL:
    for i, v in enumerate([("S2-SAMPLE1", None, None, None, None, "OWL needs a second sortie for Flight 3 next week, any PM", dt.date(YEAR, 10, 2)),
                           ("S2-SAMPLE2", None, None, None, None, "Can Friday's AM flight start 0900 instead? Ground school runs late", dt.date(YEAR, 10, 1))]):
        for c, x in enumerate(v, start=1):
            if x is not None: req.cell(6 + i, c, x)
dv = DataValidation(type="list", formula1='"AM,PM,NIGHT,WD"', allow_blank=True); dv.add(f"C6:C{5+REQ_N}"); req.add_data_validation(dv)
dv = DataValidation(type="list", formula1='"Add flight,Change,Cancel,Other"', allow_blank=True, showErrorMessage=False); dv.add(f"D6:D{5+REQ_N}"); req.add_data_validation(dv)
req.freeze_panes = "A6"; req.sheet_properties.tabColor = "D97706"

# ---------------------------------------------------------------- Syllabus Log (SUP 2): one row per syllabus attempt. The items are on the Syllabus Items sheet below;
# the dashboard's Syllabus tab counts attempts and chances from these rows. It never changes the Cat Status on the Roster.
if True:
    SYL_N = 1000
    syl = wb.create_sheet("Syllabus Log", wb.sheetnames.index("Requests") + 1)
    banner(syl, f"SYLLABUS LOG  |  {DEPT.upper()}", "One row per attempt at a syllabus item. Enter results on the dashboard (Planner > Syllabus Tracking, or the Planner cell picker), or type a row here.", 8)
    syl["A3"] = "Item Code = the Code column of the Syllabus Items sheet. Result Planned = linked to a plan cell (Date + Callsign + Shift) and not flown yet; only Pass / Fail count as attempts. Passing a stage never changes the Roster: update the Cat Status by hand once the paperwork is approved."; syl["A3"].font = f(False, SLATE500, 9, italic=True)
    header(syl, 5, 1, ["Date", "Callsign", "Item Code", "Result", "Instructor", "Remarks", "Crew Member", "Shift"])
    for col, w in zip("ABCDEFGH", [16, 12, 14, 10, 13, 44, 26, 9]): syl.column_dimensions[col].width = w
    for i in range(SYL_N):
        r = 6 + i
        for c in range(1, 9):
            cell = syl.cell(r, c); cell.border = BOX
            cell.fill = fill(SLATE100) if c == 7 else INPUT_FILL
            if c != 7: cell.font = INPUT_FONT
        syl.cell(r, 1).number_format = "ddd dd mmm yyyy"; syl.cell(r, 2).number_format = "@"; syl.cell(r, 3).number_format = "@"
        syl.cell(r, 7, f'=IF(B{r}="","",IFERROR(INDEX(Roster!$C$6:$C${5+N},MATCH(B{r},Roster!$B$6:$B${5+N},0)),"not on the Roster"))')
    D = lambda m, d: dt.date(YEAR, m, d)
    O = "S2-O-"
    SAMPLE = [  # OWL (OCU): 2 fails on Flight 2, Flight 2 planned 27 Oct; EAGLE (CAT C) part way; FALCON (CAT B) out of chances on EIS Flight 1; RAPTOR (CAT A) all passed, awaiting paperwork
        (D(7, 28), "OWL", O + "INT1", "Pass", "VIPER", "Demo flight"), (D(8, 4), "OWL", O + "INT2", "Pass", "VIPER", ""), (D(8, 6), "OWL", O + "INT3", "Pass", "VIPER", ""),
        (D(8, 11), "OWL", O + "SIM1", "Pass", "VIPER", ""), (D(8, 18), "OWL", O + "SIM2", "Pass", "HAWK", ""), (D(8, 25), "OWL", O + "SIM3", "Fail", "HAWK", "Surface picture lost after turn"),
        (D(9, 1), "OWL", O + "SIM3", "Pass", "HAWK", "2nd attempt"), (D(9, 8), "OWL", O + "SIMCHK", "Pass", "VIPER", ""), (D(9, 15), "OWL", O + "FLT1", "Pass", "VIPER", ""),
        (D(10, 8), "OWL", O + "FLT2", "Fail", "VIPER", "Weather recall, incomplete"), (D(10, 13), "OWL", O + "FLT2", "Fail", "VIPER", "Radar picture not held"),
        (D(9, 2), "EAGLE", "S2-C-TTX1", "Pass", "HAWK", ""), (D(9, 9), "EAGLE", "S2-C-SIM1", "Pass", "HAWK", ""), (D(9, 16), "EAGLE", "S2-C-SIM2", "Pass", "HAWK", ""),
        (D(10, 6), "EAGLE", "S2-C-FLT1", "Pass", "VIPER", ""), (D(10, 9), "EAGLE", "S2-C-SAL1", "Pass", "HAWK", ""),
        (D(7, 7), "FALCON", "S2-B-SIM1", "Pass", "HAWK", ""), (D(7, 14), "FALCON", "S2-B-EIS1", "Pass", "HAWK", ""), (D(7, 21), "FALCON", "S2-B-EIS2", "Fail", "VIPER", "Search pattern not completed"),
        (D(8, 12), "FALCON", "S2-B-EIS2", "Fail", "VIPER", "Comms procedure"), (D(9, 9), "FALCON", "S2-B-EIS2", "Fail", "HAWK", "Sensor handover late"),
        *[(D(5 + k // 3, 5 + 7 * (k % 3)), "RAPTOR", c, "Pass", "VIPER", "Check passed. Paperwork to the CO" if c.endswith("FLTCHK") else "") for k, c in enumerate(["S2-A-SIM1", "S2-A-SIM2", "S2-A-SIM3", "S2-A-SIMCHK", "S2-A-FLT1", "S2-A-FLT2", "S2-A-FLT3", "S2-A-FLTCHK"])],
        (D(10, 27), "OWL", O + "FLT2", "Planned", "", "", None, "NIGHT"),   # linked to OWL's NIGHT flight (NVG-2.0) on 27 Oct
    ]
    for i, v in enumerate([] if REAL else SAMPLE):
        for c, x in enumerate(v, start=1):
            if x != "": syl.cell(6 + i, c, x)
    syl.freeze_panes = "A6"
    dv = DataValidation(type="list", formula1='"Planned,Pass,Fail"', allow_blank=True, showErrorMessage=True, errorTitle="Result", error="Planned, Pass or Fail."); dv.add(f"D6:D{5+SYL_N}"); syl.add_data_validation(dv)
    dv = DataValidation(type="list", formula1=f"=Roster!$B$6:$B${5+N}", allow_blank=True, showErrorMessage=False); dv.add(f"B6:B{5+SYL_N}"); syl.add_data_validation(dv)
    syl.conditional_formatting.add(f"D6:D{5+SYL_N}", FormulaRule(formula=['$D6="Fail"'], fill=fill("FEE2E2"), font=Font(color="991B1B", bold=True)))
    syl.conditional_formatting.add(f"D6:D{5+SYL_N}", FormulaRule(formula=['$D6="Pass"'], fill=fill("DCFCE7"), font=Font(color="166534", bold=True)))
    syl.conditional_formatting.add(f"D6:D{5+SYL_N}", FormulaRule(formula=['$D6="Planned"'], fill=fill("FEF3C7"), font=Font(color="92400E", bold=True)))
    dv = DataValidation(type="list", formula1='"AM,PM,NIGHT"', allow_blank=True); dv.add(f"H6:H{5+SYL_N}"); syl.add_data_validation(dv)
    syl.sheet_properties.tabColor = "CA8A04"

# ---------------------------------------------------------------- Syllabus definition (in each department's own file): settings, stages and items.
# Each department keeps its own syllabus (SUP 1 / SUP 2 / TACCO Mission Crew, Pilot, Loadmaster). Built for every department since the 26 Oct scale-up.
SYL_DEPTS = {"SUP 1", "SUP 2", "TACCO", "Loadmaster", "Pilot"}
SYL_KIND = {"SUP 1": "Mission Crew", "SUP 2": "Mission Crew", "TACCO": "Mission Crew", "Pilot": "Pilot", "Loadmaster": "Loadmaster"}.get(DEPT, "Mission Crew")
SYL_STAGES = [("OCU", "CAT D", "Operational Conversion Unit. Passing it makes a TRAINEE a CAT D."), ("CAT C", "CAT C", ""), ("CAT B", "CAT B", ""), ("CAT A", "CAT A", ""),
              ("Instructor", "Instructor", "Gives the Instructor Role on the Roster (not a Cat Status).")]
SYL_SIM = {"Mission Crew": ["Systems familiarisation", "Sensor management", "Surface picture", "Emergency procedures"], "Pilot": ["Cockpit procedures", "Instrument flying", "Engine failures", "Crew coordination"],
           "Loadmaster": ["Cargo systems", "Load planning", "Emergency drills", "Airdrop procedures"]}
SYL_FLT = {"Mission Crew": ["Day patrol", "Radar picture", "Night patrol", "Search pattern"], "Pilot": ["General handling", "Circuits", "Night flying", "Low level"],
           "Loadmaster": ["Cargo loading", "Pax handling", "Night ops", "Airdrop"]}
SYL_PREFIX = {"Mission Crew": "MC", "Pilot": "P", "Loadmaster": "L"}
# SUP 2's real syllabus (Brandon 2026-10-24): Stage, Track, Order, Code, Description, Type (Ground / Sim / Flight). Tracks are free text, shown as tree columns in order.
SYL_REAL = {"SUP 2": [r.split(",") for r in """
OCU,Intro,1,S2-O-INT1,Demo Flight,Flight
OCU,Intro,2,S2-O-INT2,TTX 1,Ground
OCU,Intro,3,S2-O-INT3,TTX 2,Ground
OCU,Sims,1,S2-O-SIM1,SIMS 1,Sim
OCU,Sims,2,S2-O-SIM2,SIMS 2,Sim
OCU,Sims,3,S2-O-SIM3,SIMS 3,Sim
OCU,Sims,4,S2-O-SIMCHK,SIMS Check,Sim
OCU,Flight,1,S2-O-FLT1,Flight 1,Flight
OCU,Flight,2,S2-O-FLT2,Flight 2,Flight
OCU,Flight,3,S2-O-FLT3,Flight 3,Flight
OCU,Flight,4,S2-O-FLT4,Flight 4,Flight
OCU,Flight,5,S2-O-FLT5,Flight 5,Flight
OCU,Flight,6,S2-O-FLTCHK,CAT D Check,Flight
CAT C,TTX,1,S2-C-TTX1,TTX 1,Ground
CAT C,TTX,2,S2-C-TTX2,TTX 2,Ground
CAT C,Sims,1,S2-C-SIM1,SIMS 1,Sim
CAT C,Sims,2,S2-C-SIM2,SIMS 2,Sim
CAT C,Sims,3,S2-C-SIM3,SIMS 3,Sim
CAT C,Sims,4,S2-C-SIMCHK,SIMS Check,Sim
CAT C,Flight,1,S2-C-FLT1,Flight 1,Flight
CAT C,Flight,2,S2-C-FLT2,Flight 2,Flight
CAT C,Flight,3,S2-C-FLT3,Flight 3,Flight
CAT C,Flight,4,S2-C-FLTCHK,CAT C Check,Flight
CAT C,Role: SAL,1,S2-C-SAL1,SAL TTX 1,Ground
CAT C,Role: SAL,2,S2-C-SAL2,SAL Flight 1,Flight
CAT C,Role: SAL,3,S2-C-SAL3,SAL Flight 2,Flight
CAT C,Role: SAL,4,S2-C-SALCHK,SAL Check,Flight
CAT B,Sims,1,S2-B-SIM1,TACT SIMS 1,Sim
CAT B,Sims,2,S2-B-SIM2,MKZ SIMS 1,Sim
CAT B,Sims,3,S2-B-SIM3,MAS SIMS 1,Sim
CAT B,Role: EIS,1,S2-B-EIS1,EIS TTX,Ground
CAT B,Role: EIS,2,S2-B-EIS2,EIS Flight 1,Flight
CAT B,Role: EIS,3,S2-B-EISCHK,EIS Check,Flight
CAT B,Role: MS,1,S2-B-MS1,MS TTX,Ground
CAT B,Role: MS,2,S2-B-MS2,MS TACT TTX,Ground
CAT B,Role: MS,3,S2-B-MS3,MS SIMS 1,Sim
CAT B,Role: MS,4,S2-B-MS4,MS Flight 1,Flight
CAT B,Role: MS,5,S2-B-MS5,MS Flight 2,Flight
CAT B,Role: MS,6,S2-B-MSCHK,MS Check,Flight
CAT A,Sims,1,S2-A-SIM1,MAS Tact SIMS 1,Sim
CAT A,Sims,2,S2-A-SIM2,MAS Tact SIMS 2,Sim
CAT A,Sims,3,S2-A-SIM3,MAS Tact SIMS 3,Sim
CAT A,Sims,4,S2-A-SIMCHK,MAS Tact SIMS Check,Sim
CAT A,Flight,1,S2-A-FLT1,MAS Tact Flight 1,Flight
CAT A,Flight,2,S2-A-FLT2,MAS Tact Flight 2,Flight
CAT A,Flight,3,S2-A-FLT3,MAS Tact Flight 3,Flight
CAT A,Flight,4,S2-A-FLTCHK,MAS Tact Flight Check,Flight
""".strip().splitlines()]}
SYL_NPER = {"OCU": 4, "CAT C": 3, "CAT B": 3, "CAT A": 2, "Instructor": 2}
SYL_SC = {"OCU": "O", "CAT C": "C", "CAT B": "B", "CAT A": "A", "Instructor": "I"}
if DEPT in SYL_DEPTS:
    SYL_NST, SYL_NI = 20, 1000
    if DEPT in SYL_REAL: SYL_STAGES = [x for x in SYL_STAGES if x[0] != "Instructor"]
    def syl_inputs(ws, r0, n, ncol):
        for r in range(r0, r0 + n):
            for c in range(1, ncol + 1):
                cell = ws.cell(r, c); cell.fill = INPUT_FILL; cell.border = BOX; cell.font = INPUT_FONT; cell.alignment = Alignment(vertical="center", wrap_text=True)
    sset = wb.create_sheet("Syllabus Settings", wb.sheetnames.index("Syllabus Log") + 1)
    banner(sset, f"SYLLABUS SETTINGS  |  {DEPT.upper()}", "Yellow cells can be changed. Used by the dashboard's Planner > Syllabus Tracking.", 3)
    header(sset, 5, 1, ["Setting", "Value", "Meaning"])
    for col, w in zip("ABC", [32, 10, 70]): sset.column_dimensions[col].width = w
    for r, (k, v, m) in enumerate([("Default chances per item", 3, "Fails allowed on an item before the person is out of chances (used when an item's Chances cell is blank)."),
                                   ("Warn when chances left ≤", 1, "Amber warning when this many chances or fewer are left.")], start=6):
        sset.cell(r, 1, k).font = f(True); sset.cell(r, 1).border = BOX
        c = sset.cell(r, 2, v); c.fill = INPUT_FILL; c.font = f(True, "1D4ED8", 12); c.alignment = CENTER; c.border = BOX
        sset.cell(r, 3, m).font = f(False, SLATE500, 9); sset.cell(r, 3).border = BOX
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="99", allow_blank=False); dv.add(f"B{r}"); sset.add_data_validation(dv)
    sset.freeze_panes = "A6"; sset.sheet_properties.tabColor = "CA8A04"
    sstg = wb.create_sheet("Syllabus Stages", wb.sheetnames.index("Syllabus Settings") + 1)
    banner(sstg, f"SYLLABUS STAGES  |  {DEPT.upper()}", "The training path, first stage at the top. 'Gives' = the Cat Status (or Role) the stage leads to once the paperwork is approved. No gaps.", 3)
    header(sstg, 5, 1, ["Stage", "Gives", "Notes"])
    for col, w in zip("ABC", [18, 16, 70]): sstg.column_dimensions[col].width = w
    syl_inputs(sstg, 6, SYL_NST, 3)
    for i, row in enumerate(SYL_STAGES):
        for j, v in enumerate(row): sstg.cell(6 + i, 1 + j, v)
    dv = DataValidation(type="list", formula1='"TRAINEE,CAT D,CAT C,CAT B,CAT A,Instructor"', allow_blank=True); dv.add(f"B6:B{5 + SYL_NST}"); sstg.add_data_validation(dv)
    name("SylStageList", f"'Syllabus Stages'!$A$6:INDEX('Syllabus Stages'!$A$6:$A${5+SYL_NST},MAX(1,COUNTA('Syllabus Stages'!$A$6:$A${5+SYL_NST})))")
    sstg.freeze_panes = "A6"; sstg.sheet_properties.tabColor = "CA8A04"
    sit = wb.create_sheet("Syllabus Items", wb.sheetnames.index("Syllabus Stages") + 1)
    banner(sit, f"SYLLABUS ITEMS  |  {DEPT.upper()} ({SYL_KIND.upper()})", "One row per item. Track = a column of the tree (any name, items in Order). Type = Ground, Sim or Flight: the Planner cell picker offers Flight items on a flight and Sim items on a sim. Chances blank = the default." + ("" if DEPT in SYL_REAL else " Items are PLACEHOLDERS."), 8)
    sit["A3"] = "Codes must be unique: the Syllabus Log refers to items by Code. Rename a description freely, but don't change a Code that already has results."; sit["A3"].font = f(False, SLATE500, 9, italic=True)
    header(sit, 5, 1, ["Stage", "Track", "Order", "Code", "Description", "Chances", "Notes", "Type"])
    for col, w in zip("ABCDEFGH", [12, 14, 8, 15, 30, 10, 36, 10]): sit.column_dimensions[col].width = w
    syl_inputs(sit, 6, SYL_NI, 8)
    r = 6
    for st, tr, od, code, desc, ty in SYL_REAL.get(DEPT, []):
        for c, v in ((1, st), (2, tr), (3, int(od)), (4, code), (5, desc), (8, ty)): sit.cell(r, c, v)
        r += 1
    for st, _, _ in ([] if DEPT in SYL_REAL else SYL_STAGES):
        for track, lst in (("Sim", SYL_SIM[SYL_KIND]), ("Flight", SYL_FLT[SYL_KIND])):
            for k in range(SYL_NPER[st]):
                for j, v in enumerate([st, track, k + 1, f"{SYL_PREFIX[SYL_KIND]}-{SYL_SC[st]}-{'SIM' if track == 'Sim' else 'FLT'}{k + 1}", lst[k], None, None, track]): sit.cell(r, 1 + j, v)
                r += 1
        for j, v in enumerate([st, "Flight", SYL_NPER[st] + 1, f"{SYL_PREFIX[SYL_KIND]}-{SYL_SC[st]}-CHK", f"{st} check ride", 1, "Example of an item with its own chances (1).", "Flight"]): sit.cell(r, 1 + j, v)
        r += 1
    for c in (2, 3, 6, 8):
        for rr in range(6, 6 + SYL_NI): sit.cell(rr, c).alignment = CENTER
    for rr in range(6, 6 + SYL_NI): sit.cell(rr, 4).font = f(True, "1D4ED8")
    for rng, src in ((f"A6:A{5+SYL_NI}", "SylStageList"), (f"H6:H{5+SYL_NI}", '"Ground,Sim,Flight"')):
        dv = DataValidation(type="list", formula1=src, allow_blank=True); dv.add(rng); sit.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="99", allow_blank=True); dv.add(f"F6:F{5+SYL_NI}"); sit.add_data_validation(dv)
    sit.freeze_panes = "A6"; sit.sheet_properties.tabColor = "CA8A04"

# ---------------------------------------------------------------- Personnel Info: one row per roster slot, same row numbers as the Roster
if True:
    pers = wb.create_sheet("Personnel Info", wb.sheetnames.index("Roster") + 1)
    PCOLS = [("#", 5), ("Callsign", 12), ("Crew Member", 28), ("DOB", 13), ("Age", 6), ("NRIC", 13), ("Pers ID", 13), ("Passport No", 14), ("Passport Issue", 14),
             ("Passport Expiry", 14), ("Passport", 12), ("Contact", 14), ("Vehicle Plate", 13), ("Address", 40), ("Blood Type", 10), ("NOK", 24), ("NOK Contact", 14), ("Allergies", 26)]
    PL = len(PCOLS)
    banner(pers, f"PERSONNEL INFO  |  {DEPT.upper()}", "Restricted: personal data. Callsign and name come from the Roster (same rows), so add or change crew on the Roster and fill in their details here.", PL)
    pers["A3"] = "Grey columns fill in by themselves. Passport turns amber when it expires within 6 months and red once expired."; pers["A3"].font = f(False, SLATE500, 9, italic=True)
    header(pers, 5, 1, [c for c, _ in PCOLS])
    pers.row_dimensions[5].height = 22
    DATEC, TEXTC = (4, 9, 10), (6, 7, 8, 12, 13, 17)
    for i in range(N):
        r = 6 + i
        pers.cell(r, 1, i + 1).font = f(color=SLATE400); pers.cell(r, 1).alignment = CENTER
        pers.cell(r, 2, f'=IF(Roster!B{r}="","",Roster!B{r})').font = f(True, SLATE900)
        pers.cell(r, 3, f'=IF(Roster!C{r}="","",Roster!C{r})')
        pers.cell(r, 5, f'=IF(D{r}="","",DATEDIF(D{r},TODAY(),"y"))').alignment = CENTER
        pers.cell(r, 11, f'=IF(J{r}="","",IF(J{r}<TODAY(),"EXPIRED",IF(J{r}<EDATE(TODAY(),6),"< 6 months","Valid")))').alignment = CENTER
        for c in range(1, PL + 1):
            cell = pers.cell(r, c); cell.border = BOX
            auto = c in (2, 3, 5, 11)
            cell.fill = fill(SANS_FILL) if is_sans(i) else fill(SLATE100) if auto or c == 1 else INPUT_FILL
            if not auto and c > 1: cell.font = INPUT_FONT
            if c in DATEC: cell.number_format = "dd mmm yyyy"; cell.alignment = CENTER
            if c in TEXTC: cell.number_format = "@"
            if c == 15: cell.alignment = CENTER
            if c == 14: cell.alignment = Alignment(vertical="center", wrap_text=True)
    for c in range(1, PL + 1):
        pers.cell(SANS_ROW0, c).border = Border(left=thin, right=thin, top=Side(style="thick", color="7E22CE"), bottom=thin)
        pers.column_dimensions[CL(c)].width = PCOLS[c - 1][1]
    pers.freeze_panes = "D6"
    dv = DataValidation(type="list", formula1='"A+,A-,B+,B-,AB+,AB-,O+,O-"', allow_blank=True); dv.add(f"O6:O{5+N}"); pers.add_data_validation(dv)
    dv = DataValidation(type="date", operator="between", formula1="DATE(1940,1,1)", formula2="TODAY()", allow_blank=True, showErrorMessage=True, errorTitle="Date of birth", error="Type a date, e.g. 14/03/1990.")
    dv.add(f"D6:D{5+N}"); pers.add_data_validation(dv)
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)", allow_blank=True, showErrorMessage=True, errorTitle="Passport date", error="Type a date, e.g. 01/06/2030.")
    dv.add(f"I6:J{5+N}"); pers.add_data_validation(dv)
    pers.conditional_formatting.add(f"K6:K{5+N}", FormulaRule(formula=['$K6="EXPIRED"'], fill=fill("DC2626"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    pers.conditional_formatting.add(f"K6:K{5+N}", FormulaRule(formula=['$K6="< 6 months"'], fill=fill("F59E0B"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    pers.conditional_formatting.add(f"K6:K{5+N}", FormulaRule(formula=['$K6="Valid"'], fill=fill("DCFCE7"), font=Font(color="166534", bold=True), stopIfTrue=True))
    pers.conditional_formatting.add(f"R6:R{5+N}", FormulaRule(formula=['$R6<>""'], font=Font(color="B91C1C", bold=True)))
    pers.conditional_formatting.add(f"B6:R{5+N}", FormulaRule(formula=['INDEX(CrewStatus,ROW()-5)="Inactive"'], font=Font(color=SLATE400, italic=True, strike=True)))
    cat_cf(pers, f"C6:C{5+N}", "ROW()-5")
    pers.cell(SANS_ROW0, PL + 2, "SANS / NSmen from here").font = f(True, "7E22CE", 9)
    pers.sheet_properties.tabColor = "B91C1C"
    if not NOFILTER: pers.auto_filter.ref = f"A5:{CL(PL)}{5+N}"

# ---------------------------------------------------------------- Activities
banner(acts, "ACTIVITY DROPDOWN CATALOG", "Flights programme first (category Flight: copied from the Master Planner's Flight Codes in the dashboard), then the ground programme (everything else). Keep codes in one block with no gaps.", 10)
header(acts, 5, 1, ["Code", "Description", "Category", "Flight Hrs", "Sorties"])
# flights programme first (the Master Planner's Flight Codes, copied in by the dashboard), then the ground programme
catalog = [("FLT-1.5", "Routine Sortie (1.5h)", "Flight", 1.5), ("FLT-2.0", "Standard Sortie (2.0h)", "Flight", 2.0),
           ("FLT-3.0", "Extended Mission (3.0h)", "Flight", 3.0), ("NVG-2.0", "Night Flying Sortie (2.0h)", "Flight", 2.0),
           ("2xFLT-1.5", "Two routine sorties in one shift", "Flight", 3.0, 2), ("2xFLT-2.0", "Two standard sorties in one shift", "Flight", 4.0, 2),
           ("MSN-3.0", "Mission sortie: Pilot, TACCO, SUP 2 (3.0h)", "Flight", 3.0), ("PLT-TRG", "Pilot training sortie (1.5h)", "Flight", 1.5),
           ("EIS", "EIS training sortie (2.0h)", "Flight", 2.0), ("MAS TRG", "MAS tactical training (3.0h)", "Flight", 3.0),
           # ground programme: Brandon's list of 2026-10-21 (Activities_List_to_fill.xlsx)
           ("SIM", "Simulator Training", "Sim", 0), ("SIMCEN", "Simulator Training (SIMCEN)", "Sim", 0),
           ("BU", "Quick Reaction Standby", "Ground Duty", 0),
           ("P1", "P1 Crew", "Ground Duty", 0), ("P2", "P2 Crew", "Ground Duty", 0),
           *([("SXO", "Squadron Executive Officer", "Squadron Duty", 0)] if DEPT == "Pilot" else []),
           *([("OPS O", "Operations Officer", "Squadron Duty", 0)] if DEPT in ("Pilot", "TACCO") else []),
           *([("SDO", "Squadron Duty Orderly", "Squadron Duty", 0)] if DOES_SDO else []),
           ("BDO", "Base Duty Officer", "Squadron Duty", 0),
           ("Meeting", "Meeting", "Admin", 0), ("Outstation", "Out of office for work purpose", "Admin", 0),
           ("Detachment", "Detachment", "Oversea Duty", 0),
           ("MC", "Medical Leave", "Medical", 0), ("Medical Appt", "Medical Appointment", "Medical", 0),
           ("LL", "Local Leave", "Leave", 0), ("OL", "Oversea Leave", "Leave", 0), ("OIL", "OIL", "Off", 0)]
if REAL:   # the Master Planner's flight codes (all of them), then this department's activities from the data input
    catalog = [(c["code"], c["desc"], "Flight", c["hrs"]) + ((c["sorties"],) if c["sorties"] != 1 else ()) for c in realdata.FLIGHT_CODES] + \
              [(a["code"], a["desc"], a["cat"], 0) for a in realdata.acts_for(DEPT)]
for i in range(NA):
    r = 6 + i
    for c in range(1, 6):
        cell = acts.cell(r, c); cell.fill = INPUT_FILL; cell.font = INPUT_FONT; cell.border = BOX
    acts.cell(r, 4).number_format = "0.0;-0.0;-"
    if i < len(catalog):
        for j, v in enumerate(catalog[i]):
            acts.cell(r, 1 + j, v)
header(acts, 5, 7, ["Category"])
for i, (cat, bg, fg, it) in enumerate(CATS):
    c = acts.cell(6 + i, 7, cat); c.fill = fill(bg); c.font = f(True, fg, italic=it); c.border = BOX
acts.cell(6 + len(CATS) + 1, 7, "You can rename these; don't add more.").font = f(italic=True, color=SLATE500, size=8)
header(acts, 5, 9, ["Public Holiday", "Name"])
# Singapore public holidays as gazetted by MOM (in-lieu Mondays included); the plan year and the next year are listed
SG_HOLS = {
    2026: [("2026-01-01", "New Year's Day"), ("2026-02-17", "Chinese New Year"), ("2026-02-18", "Chinese New Year"), ("2026-03-21", "Hari Raya Puasa"),
           ("2026-04-03", "Good Friday"), ("2026-05-01", "Labour Day"), ("2026-05-27", "Hari Raya Haji"), ("2026-05-31", "Vesak Day"),
           ("2026-06-01", "Vesak Day (in lieu)"), ("2026-08-09", "National Day"), ("2026-08-10", "National Day (in lieu)"), ("2026-11-08", "Deepavali"),
           ("2026-11-09", "Deepavali (in lieu)"), ("2026-12-25", "Christmas Day")],
    2027: [("2027-01-01", "New Year's Day"), ("2027-02-06", "Chinese New Year"), ("2027-02-07", "Chinese New Year"), ("2027-02-08", "Chinese New Year (in lieu)"),
           ("2027-03-10", "Hari Raya Puasa"), ("2027-03-26", "Good Friday"), ("2027-05-01", "Labour Day"), ("2027-05-17", "Hari Raya Haji"),
           ("2027-05-20", "Vesak Day"), ("2027-08-09", "National Day"), ("2027-10-28", "Deepavali"), ("2027-12-25", "Christmas Day")],
}
hols = [(dt.date.fromisoformat(d), n) for y in (YEAR, YEAR + 1) for d, n in SG_HOLS.get(y, [])][:30]
if REAL and realdata.HOLS: hols = realdata.HOLS[:30]
if not SG_HOLS.get(YEAR): hols = [(dt.date(YEAR, 1, 1), "New Year's Day"), (dt.date(YEAR, 5, 1), "Labour Day"), (dt.date(YEAR, 12, 25), "Christmas Day")] + hols
for i in range(30):
    r = 6 + i
    for c in (9, 10):
        cell = acts.cell(r, c); cell.fill = INPUT_FILL; cell.font = INPUT_FONT; cell.border = BOX
    acts.cell(r, 9).number_format = "dd-mmm-yy"
    if i < len(hols):
        acts.cell(r, 9, hols[i][0]); acts.cell(r, 10, hols[i][1])
for col, w in zip("ABCDEFGHIJK", [12, 30, 15, 11, 30, 3, 15, 3, 15, 22, 15]):
    acts.column_dimensions[col].width = w
# K: syllabus tag per flight code (copied from the Master Planner's Flight Codes H with the flight codes)
header(acts, 5, 11, ["Syllabus tag"])
for i in range(NA):
    cell = acts.cell(6 + i, 11); cell.fill = INPUT_FILL; cell.font = INPUT_FONT; cell.border = BOX
    if i < len(catalog): cell.value = ({c["code"]: c["tag"] for c in realdata.FLIGHT_CODES} if REAL else {"EIS": "EIS", "MAS TRG": "MAS Tact"}).get(catalog[i][0]) or None
acts.freeze_panes = "A6"
dv = DataValidation(type="list", formula1="Categories", allow_blank=True); dv.add(f"C6:C{5+NA}"); acts.add_data_validation(dv)
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True); dv.add(f"D6:D{5+NA}"); acts.add_data_validation(dv)
dv = DataValidation(type="custom", formula1=f"COUNTIF($A$6:$A${5+NA},A6)=1", allow_blank=True, showErrorMessage=True,
                    errorTitle="Duplicate code", error="That code already exists."); dv.add(f"A6:A{5+NA}"); acts.add_data_validation(dv)
apply_cat_cf(acts, f"A6:A{5+NA}", "A6")
acts.cell(4, 5, "Blank = 1").font = f(italic=True, color=SLATE500, size=8)
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="9", allow_blank=True); dv.add(f"E6:E{5+NA}"); acts.add_data_validation(dv)
header(acts, 5, 12, ["Sorties (calc)", "Flight #", "Flight list", "WD #", "WD list"])
for i in range(NA):
    r = 6 + i
    acts.cell(r, 12, f'=IF($A{r}="","",IF($E{r}="",1,$E{r}))')
    acts.cell(r, 13, f'=IF(AND($A{r}<>"",$C{r}=$G$6),COUNTIFS($A$6:$A{r},"<>",$C$6:$C{r},$G$6),"")')
    acts.cell(r, 14, f'=IFERROR(INDEX($A$6:$A${5+NA},MATCH({i + 1},$M$6:$M${5+NA},0)),"")')
    # the WD row takes ground and admin items only
    acts.cell(r, 15, f'=IF(AND($A{r}<>"",OR($C{r}={catref("Ground Duty")},$C{r}={catref("Admin")})),COUNTIFS($A$6:$A{r},"<>",$C$6:$C{r},{catref("Ground Duty")})+COUNTIFS($A$6:$A{r},"<>",$C$6:$C{r},{catref("Admin")}),"")')
    acts.cell(r, 16, f'=IFERROR(INDEX($A$6:$A${5+NA},MATCH({i + 1},$O$6:$O${5+NA},0)),"")')
for col in "LMNOP":
    acts.column_dimensions[col].hidden = True
name("ActSort", f"Activities!$L$6:$L${5+NA}")
name("FlightCodes", f"Activities!$N$6:$N${5+NA}" if STATICNAMES else f"Activities!$N$6:INDEX(Activities!$N$6:$N${5+NA},MAX(1,COUNTIF(Activities!$C$6:$C${5+NA},Activities!$G$6)))")
name("WdCodes", f"Activities!$P$6:$P${5+NA}" if STATICNAMES else f"Activities!$P$6:INDEX(Activities!$P$6:$P${5+NA},MAX(1,COUNTIF(Activities!$C$6:$C${5+NA},{catref('Ground Duty')})+COUNTIF(Activities!$C$6:$C${5+NA},{catref('Admin')})))")
name("ActCodesAll", f"Activities!$A$6:$A${5+NA}")
name("ActCat", f"Activities!$C$6:$C${5+NA}")
name("ActHrs", f"Activities!$D$6:$D${5+NA}")
name("ActCodes", f"Activities!$A$6:$A${5+NA}" if STATICNAMES else f"Activities!$A$6:INDEX(Activities!$A$6:$A${5+NA},MAX(1,COUNTA(Activities!$A$6:$A${5+NA})))")
name("Categories", f"Activities!$G$6:$G${5+len(CATS)}")
name("Holidays", "Activities!$I$6:$I$35")

# ---------------------------------------------------------------- Currency checks (SUP 2 only for now)
# Mission type per activity (Activities F), a rule table anyone can extend (Currency Rules) and a crew x rule sheet (Currency)
# holding the "last flown before the planner" baseline dates. The dashboard works out each status and writes it in on Save.
HAS_CUR = True
CUR_RULES, CUR_ROW0 = 12, 6
if HAS_CUR:
    header(acts, 5, 6, ["Mission Type"])
    mtypes = {"FLT-1.5": "GEN", "FLT-2.0": "GEN", "FLT-3.0": "GEN", "NVG-2.0": "NIGHT", "2xFLT-1.5": "GEN", "2xFLT-2.0": "GEN", "MSN-3.0": "MSN", "PLT-TRG": "TRG", "EIS": "TRG", "MAS TRG": "TRG"}
    for i in range(NA):
        r = 6 + i; cell = acts.cell(r, 6); cell.fill = INPUT_FILL; cell.font = INPUT_FONT; cell.border = BOX
        if REAL: mtypes = {c["code"]: c["mtype"] for c in realdata.FLIGHT_CODES if c["mtype"]}
        if i < len(catalog) and catalog[i][0] in mtypes: cell.value = mtypes[catalog[i][0]]
    acts.column_dimensions["F"].width = 13
    acts.cell(4, 6, "for currency rules").font = f(italic=True, color=SLATE500, size=8)

    rules = wb.create_sheet("Currency Rules", wb.sheetnames.index("Activities"))
    rules.sheet_properties.tabColor = "7C3AED"
    banner(rules, "CURRENCY RULES", "One rule per row. A person is flagged when they haven't done enough qualifying flights within the window. Rules apply to pilots and non-pilots alike.", 9)
    header(rules, 5, 1, ["Rule", "Applies to", "Who", "Counts", "Which", "Within (days)", "At least (flights)", "Warn (days before)", "Notes"])
    samples = [("General flying", "All", "", "Any flight", "", 45, 1, 7, "Flagged after 45 calendar days without flying"),
               ("Mission currency", "All", "", "Mission type", "MSN", 180, 1, 14, "Flagged after 180 days without a mission sortie (Mission Type MSN on Activities)"),
               ("Night flying", "Role", "Instructor", "Mission type", "NIGHT", 90, 2, 14, "Example of a count rule: at least 2 night sorties in 90 days")]
    if REAL: samples = samples[:1]   # the real codes have no MSN / NIGHT mission types: keep the general rule, the squadron adds its own
    for k in range(CUR_RULES):
        r = CUR_ROW0 + k
        for c in range(1, 10):
            cell = rules.cell(r, c); cell.fill = INPUT_FILL; cell.font = INPUT_FONT; cell.border = BOX; cell.alignment = Alignment(vertical="center", wrap_text=c in (3, 5, 9))
        if k < len(samples):
            for j, v in enumerate(samples[k]): rules.cell(r, 1 + j, v if v != "" else None)
        rules.row_dimensions[r].height = 30
    dv = DataValidation(type="list", formula1='"All,Role,Cat Status,Status,Callsigns"', allow_blank=True); dv.add(f"B6:B{5+CUR_RULES}"); rules.add_data_validation(dv)
    dv = DataValidation(type="list", formula1='"Any flight,Mission type,Flight codes"', allow_blank=True); dv.add(f"D6:D{5+CUR_RULES}"); rules.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3650", allow_blank=True); dv.add(f"F6:F{5+CUR_RULES}"); rules.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="99", allow_blank=True); dv.add(f"G6:H{5+CUR_RULES}"); rules.add_data_validation(dv)
    for col, w in zip("ABCDEFGHI", [22, 13, 22, 15, 22, 11, 11, 11, 55]): rules.column_dimensions[col].width = w
    tips = ["How to write a rule",
            "Applies to: All, or Role / Cat Status / Status / Callsigns with the values in Who, comma separated (e.g. Instructor, Senior CAT A  or  VIPER, HAWK).",
            "Counts: Any flight (category Flight on Activities), Mission type (Activities column F, e.g. MSN) or Flight codes (e.g. NVG-2.0, MSN-3.0).",
            "At least: blank = 1. A 2xFLT code counts as 2 flights.",
            "Warn: the person shows amber this many days before the rule lapses.",
            f"Up to {CUR_RULES} rules. Past dates on the month tabs count as flown; change or clear a cell if the flight was cancelled.",
            "Flying before this file starts: enter each person's last date on the Currency sheet."]
    for i, t in enumerate(tips):
        c = rules.cell(CUR_ROW0 + CUR_RULES + 2 + i, 1, t); c.font = f(i == 0, SLATE700 if i else "5B21B6", 10 if i == 0 else 9)
    rules.freeze_panes = "A6"

    cur = wb.create_sheet("Currency", wb.sheetnames.index("Currency Rules"))
    cur.sheet_properties.tabColor = "7C3AED"
    LCC = 3 + 2 * CUR_RULES
    banner(cur, "CURRENCY", "Status is worked out by the dashboard and written here when you Save to Excel. Enter each person's last qualifying flight before this file started in the 'Before' columns.", LCC)
    cur["A3"] = "As of"; cur["A3"].font = f(True, SLATE700, 9); cur["B3"].number_format = "dd-mmm-yy"; cur["B3"].font = f(True, "5B21B6", 10)
    cur["C3"] = "(date of the last dashboard save)"; cur["C3"].font = f(False, SLATE500, 8, italic=True)
    for c, lab in ((1, "Callsign"), (2, "Name"), (3, "Flags")):
        cur.merge_cells(start_row=4, start_column=c, end_row=5, end_column=c)
        h = cur.cell(4, c, lab); h.font = f(True, "FFFFFF", 10); h.fill = fill(SLATE900); h.alignment = CENTER; cur.cell(5, c).fill = fill(SLATE900)
    for k in range(CUR_RULES):
        c0 = 4 + 2 * k; rr = CUR_ROW0 + k
        cur.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 1)
        h = cur.cell(4, c0, f"=IF('Currency Rules'!$A${rr}=\"\",\"Rule {k + 1}\",'Currency Rules'!$A${rr})"); h.font = f(True, "FFFFFF", 10); h.fill = fill("5B21B6"); h.alignment = CENTER
        cur.cell(4, c0 + 1).fill = fill("5B21B6")
        for j, lab in enumerate(("Before", "Status")):
            x = cur.cell(5, c0 + j, lab); x.font = f(True, "5B21B6", 8); x.fill = fill("EDE9FE"); x.alignment = CENTER; x.border = BOX
        cur.column_dimensions[CL(c0)].width = 10; cur.column_dimensions[CL(c0 + 1)].width = 24
    base = {0: {1: "2026-05-02"}, 1: {1: "2026-06-20", 2: "2026-09-01"}, 2: {0: "2026-09-10", 1: "2026-08-15"}, 3: {0: "2026-09-12"},
            4: {1: "2026-05-20"}, REG_N: {0: "2026-08-20"}, REG_N + 1: {0: "2026-09-25", 1: "2026-07-07"}}
    for i in range(N):
        r = CUR_ROW0 + i; rr = 6 + i
        a = cur.cell(r, 1, f'=IF(Roster!$C${rr}="","",IF(Roster!$B${rr}="","#"&{i + 1},Roster!$B${rr}&""))'); a.font = f(True, SLATE700, 9); a.alignment = CENTER
        b = cur.cell(r, 2, f'=IF(Roster!$C${rr}="","",Roster!$C${rr})'); b.font = f(False, SLATE800, 10)
        cc = cur.cell(r, 3); cc.alignment = CENTER; cc.font = f(True, "B91C1C", 10)
        for k in range(CUR_RULES):
            c0 = 4 + 2 * k
            x = cur.cell(r, c0); x.fill = INPUT_FILL; x.font = INPUT_FONT; x.number_format = "dd-mmm-yy"; x.border = BOX; x.alignment = CENTER
            if k in base.get(i, {}) and not REAL: x.value = dt.date.fromisoformat(base[i][k])
            y = cur.cell(r, c0 + 1); y.border = BOX; y.font = f(False, SLATE700, 9); y.alignment = Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True)
        for c in range(1, 4): cur.cell(r, c).border = BOX
        if is_sans(i):
            for c in (1, 2): cur.cell(r, c).fill = fill(SANS_FILL)
    last = CUR_ROW0 + N - 1
    srng = f"E{CUR_ROW0}:{CL(LCC)}{last}"
    for word, bg, fg in (("LAPSED", "DC2626", "FFFFFF"), ("NO RECORD", "DC2626", "FFFFFF"), ("DUE", "F59E0B", "111827"), ("OK", "DCFCE7", "166534")):
        cur.conditional_formatting.add(srng, FormulaRule(formula=[f'AND(MOD(COLUMN(E{CUR_ROW0}),2)=1,LEFT(E{CUR_ROW0},{len(word)})="{word}")'], fill=fill(bg), font=Font(color=fg, bold=True), stopIfTrue=True))
    cur.conditional_formatting.add(f"C{CUR_ROW0}:C{last}", FormulaRule(formula=[f'N(C{CUR_ROW0})>0'], fill=fill("FEE2E2"), font=Font(color="B91C1C", bold=True)))
    cur.column_dimensions["A"].width = 11; cur.column_dimensions["B"].width = 24; cur.column_dimensions["C"].width = 7
    cur.freeze_panes = "D6"
    nonblank_filter(cur, f"A5:{CL(LCC)}{last}", 0)

# ---------------------------------------------------------------- SDO Plan (year calendar of this department's SDO days)
# month block mi sits at block-row mi // 3 and block-column mi % 3; each block: title, weekday header, 6 weeks of (date row, SDO row)
SDO_BW, SDO_BH = 8, 16
def sdo_block(mi):
    return 1 + (mi % 3) * SDO_BW, 6 + (mi // 3) * SDO_BH   # first column, title row
SDO_LAST = 3 * SDO_BW - 1
SC0 = SDO_LAST + 3          # summary table (one gap column after the calendars)
def build_plan(sdo, CODE, LONG, DOES, COLOR):
    banner(sdo, f"{CODE} PLANNING CALENDAR  |  {DEPT.upper()}", (f"{LONG} ({CODE}) days for {DEPT}. Copy them from the Duty Roster with the dashboard (Planner tab), or pick your department on each day. The month tabs and the table on the right fill in automatically." if HAS_DUTY else "Pick your department in the dropdown on each day your department has SDO duty. The month tabs and the table on the right fill in automatically."), SC0 + 2)
    sdo["A4"] = f"Other departments' {CODE} days are kept in their own files. Open them in the dashboard to see the combined calendar."
    if not DOES:
        sdo["A4"] = f"{DEPT} does not do SDO duty, so this calendar stays empty. Only SUP 1, SUP 2 and Loadmaster are rostered for SDO."
    sdo["A4"].font = f(italic=True, color=SLATE500, size=9)
    dv_sdo = DataValidation(type="list", formula1="ThisDept", allow_blank=True, showErrorMessage=True, errorStyle="stop",
                            errorTitle=f"{CODE} day", error="Pick your department from the dropdown, or leave the day empty.")
    DOW = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    for mi in range(13):   # 12 months + next January (block 13 sits under January, rows 70-85)
        c0, r0 = sdo_block(mi)
        first = f"DATE(PlanYear,{mi + 1},1)"
        sdo.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 6)
        t = sdo.cell(r0, c0, f'=UPPER(TEXT({first},"mmmm yyyy"))'); t.font = f(True, "FFFFFF", 11); t.alignment = CENTER
        for c in range(c0, c0 + 7):
            sdo.cell(r0, c).fill = fill(SLATE900)
            h = sdo.cell(r0 + 1, c, DOW[c - c0]); h.font = f(True, AMBER300 if c - c0 >= 5 else "FFFFFF", 9); h.fill = fill(SLATE700 if c - c0 >= 5 else SLATE800); h.alignment = CENTER
        sdo.row_dimensions[r0].height = 20; sdo.row_dimensions[r0 + 1].height = 15
        for w in range(6):
            rd, rp = r0 + 2 + 2 * w, r0 + 3 + 2 * w
            for wd in range(7):
                c = c0 + wd; n = w * 7 + wd + 1
                dcell = sdo.cell(rd, c, f'=IF(AND({n}-(WEEKDAY({first},2)-1)>=1,{n}-(WEEKDAY({first},2)-1)<=DAY(EOMONTH({first},0))),{first}+{n}-WEEKDAY({first},2),"")')
                dcell.number_format = "d"; dcell.font = f(True, SLATE500, 9); dcell.alignment = Alignment(horizontal="left", vertical="top", indent=0)
                dcell.border = Border(left=thin, right=thin, top=thin)
                pcell = sdo.cell(rp, c); pcell.fill = INPUT_FILL; pcell.font = f(True, "1D4ED8", 10)
                pcell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True); pcell.border = Border(left=thin, right=thin, bottom=thin)
            sdo.row_dimensions[rd].height = 14; sdo.row_dimensions[rp].height = 24
            L0, L6 = CL(c0), CL(c0 + 6)
            rng = f"{L0}{rp}:{L6}{rp}"
            if DOES:
                dv_sdo.add(rng)
            sdo.conditional_formatting.add(f"{L0}{rd}:{L6}{rp}", FormulaRule(formula=[f'{L0}${rd}=""'], fill=fill(SLATE200), stopIfTrue=True))
            sdo.conditional_formatting.add(rng, FormulaRule(formula=[f'{L0}{rp}<>""'], fill=fill("FECDD3"), font=Font(color="881337", bold=True), stopIfTrue=True))
            sdo.conditional_formatting.add(f"{L0}{rd}:{L6}{rd}", FormulaRule(formula=[f'COUNTIF(Holidays,{L0}{rd})>0'], fill=fill("FFE4E6"), font=Font(color="9F1239", bold=True)))
        for c in range(c0, c0 + 7):
            sdo.column_dimensions[CL(c)].width = 12
        if mi % 3 < 2:
            sdo.column_dimensions[CL(c0 + 7)].width = 2
        if mi % 3 == 0:
            sdo.row_dimensions[r0 + SDO_BH - 2].height = 10; sdo.row_dimensions[r0 + SDO_BH - 1].height = 10
    sdo.add_data_validation(dv_sdo)
    sdo.column_dimensions[CL(SDO_LAST + 1)].width = 2; sdo.column_dimensions[CL(SDO_LAST + 2)].width = 2
    sdo.cell(4, SC0, f'="{CODE} DATES  |  "&UPPER(ThisDept)').font = f(True, SLATE900, 11)
    header(sdo, 5, SC0, ["Month", f"{CODE} dates", "Days"])
    for k, w in enumerate([11, 26, 7]):
        sdo.column_dimensions[CL(SC0 + k)].width = w
    for mi, mon in enumerate(MONTHS13):
        r = 6 + mi + (mi == 12)   # next January goes under the Year total
        lk = sdo.cell(r, SC0, f'=HYPERLINK("#\'{mon}\'!D{G0}","{MONTHS_FULL[mi] if mi < 12 else "January (next year)"}")')
        lk.font = Font(name=FONT, color="1D4ED8", underline="single", size=9); lk.border = BOX; lk.alignment = Alignment(vertical="top")
        terms = "&".join(f'IF(ISNUMBER(SEARCH("{CODE}",\'{mon}\'!{CL(D1 + d)}${SDO_DEPT_ROW})),", "&{d + 1},"")' for d in range(31))
        c = sdo.cell(r, SC0 + 1, f"=MID({terms},3,200)")
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True); c.font = f(False, SLATE800, 10); c.border = BOX
        c = sdo.cell(r, SC0 + 2, f"=COUNTIF('{mon}'!$D${SDO_DEPT_ROW}:$AH${SDO_DEPT_ROW},\"*{CODE}*\")")
        c.alignment = CENTER; c.font = f(True, SLATE900, 10); c.border = BOX; c.number_format = "0;-0;-"
        sdo.row_dimensions[r].height = max(sdo.row_dimensions[r].height or 15, 26)
    sdo.cell(18, SC0, "Year").font = f(True); sdo.cell(18, SC0).fill = fill(SLATE100); sdo.cell(18, SC0).border = BOX
    c = sdo.cell(18, SC0 + 2, f"=SUM({CL(SC0 + 2)}6:{CL(SC0 + 2)}17)"); c.font = f(True); c.alignment = CENTER; c.border = BOX; c.fill = fill(SLATE100)
    sdo.cell(18, SC0 + 1).fill = fill(SLATE100); sdo.cell(18, SC0 + 1).border = BOX
    sdo.freeze_panes = "A6"
    sdo.sheet_properties.tabColor = COLOR
build_plan(sdo, "SDO", "Squadron Duty Orderly", DOES_SDO, "BE123C")
if HAS_DUTY:
    build_plan(bdo, "BDO", "Base Duty Officer", True, "1D4ED8")

# ---------------------------------------------------------------- Month tabs
seeds = seedplan.FIXED
# sample SDO calendar for October: day -> dept (8 Oct is given to SUP 1 but nobody from SUP 1 is on SDO, to show the warning)
sdo_seed = {1: "SUP 1", 2: "SUP 2", 3: "SUP 2", 4: "SUP 2", 5: "Loadmaster", 6: "SUP 1", 7: "Loadmaster", 8: "SUP 1", 9: "SUP 2"}
SHIFTS = ["AM", "PM", "NIGHT"]
SHIFT_STYLE = {"AM": ("F0F9FF", "075985"), "PM": ("FFFBEB", "92400E"), "NIGHT": ("1E1B4B", "C7D2FE"), "WD": ("CCFBF1", "115E59")}
ROW_LABELS = SHIFTS + ["WD"]
WD_TINT = "F0FDFA"
KPI = [  # label, formula builder, accent, numfmt
    ("Total Flights", lambda: f"=SUM(AI{G0}:AI{GL})", "1D4ED8", "0"),
    ("Total Flying Hrs", lambda: f"=SUM(AJ{G0}:AJ{GL})", "0F766E", "0.0"),
    ("Ground Duties", lambda: f"=SUM(AK{G0}:AK{GL})", "B45309", "0"),
    ("SDO Planned", lambda: f"=SUM(AL{G0}:AL{GL})", "BE123C", "0"),
    ("Active Crew", lambda: '=(COUNTIFS(CrewName,"<>")-COUNTIFS(CrewName,"<>",CrewStatus,"Inactive"))', SLATE700, "0"),
    ("SDO Days Uncovered", lambda: f'=SUMPRODUCT(($D${SDO_DEPT_ROW}:$AH${SDO_DEPT_ROW}<>"")*($D${SDO_ROW}:$AH${SDO_ROW}=0))', "DC2626", "0"),
    ("Uncrewed Shifts", lambda: "=" + "+".join(f'SUMPRODUCT(($D$5:$AH$5<>"")*((($D${CODE_ROW(s, 0)}:$AH${CODE_ROW(s, 0)}<>"")+($D${CODE_ROW(s, 1)}:$AH${CODE_ROW(s, 1)}<>""))>0)*($D${S0 + 1 + s}:$AH${S0 + 1 + s}=0))' for s in range(3)), "B91C1C", "0"),
]
S0 = GL + 2
SDO_ROW = S0 + 6
SUMMARY = ["Flights AM", "Flights PM", "Flights NIGHT", "Total flights", "Fly hrs", "SDO on duty",
           "Gnd duty (crew)", "On leave (crew)", "Available crew"]
SUM_ROW = {lab: S0 + 1 + i for i, lab in enumerate(SUMMARY)}
assert SUM_ROW["SDO on duty"] == SDO_ROW

for m, mon in enumerate(MONTHS13, start=1):
    ws = wb.create_sheet(mon)
    ws.sheet_properties.tabColor = "7C3AED" if m == 13 else "1E40AF" if m != 10 else "DC2626"
    banner(ws, "", f"{DEPT}  •  3-Shift Daily Matrix  •  Pick each AM / PM / NIGHT cell from the dropdown  •  Totals, flying hours, ground duty and SDO calculate automatically", LASTC)
    ws["A1"] = f'=UPPER(ThisDept)&"  |  LODESTAR  |  AM / PM / NIGHT FLIGHT & OPS PLANNING  |  "&UPPER(TEXT(A3,"mmmm yyyy"))'
    ws["A3"] = f"=DATE(PlanYear,{m},1)"; ws["A3"].number_format = ";;;"
    ws["B3"] = "=DAY(EOMONTH(A3,0))"; ws["B3"].number_format = ";;;"
    ws.row_dimensions[3].height = 6
    for k, (lab, fx, accent, nf) in enumerate(KPI):
        ws.cell(3, D1 + k, fx()).number_format = ";;;"
    for k, (lab, cat, mode, col) in enumerate(METRICS):
        if cat: ws.cell(3, MC0 + k, f"={catref(cat)}").number_format = ";;;"

    # matrix header (rows 4-5)
    for c, lab in ((1, "Callsign"), (2, "Crew Member"), (3, "Shift")):
        ws.merge_cells(start_row=4, start_column=c, end_row=5, end_column=c)
        h = ws.cell(4, c, lab); h.font = f(True, "FFFFFF", 10); h.fill = fill(SLATE900); h.alignment = CENTER
        ws.cell(5, c).fill = fill(SLATE900)
    for d in range(1, 32):
        col = D1 + d - 1; L = CL(col)
        n = ws.cell(5, col, f'=IF({d}<=$B$3,$A$3+{d - 1},"")'); n.number_format = "dd"
        w = ws.cell(4, col, f'=IF({L}5="","",{L}5)'); w.number_format = "ddd"
        n.font = f(True, "FFFFFF", 11); w.font = f(False, SLATE400, 8)
        for x in (n, w):
            x.fill = fill(SLATE800); x.alignment = CENTER
    for k, (lab, cat, mode, col) in enumerate(METRICS):
        ws.merge_cells(start_row=4, start_column=MC0 + k, end_row=5, end_column=MC0 + k)
        h = ws.cell(4, MC0 + k, lab); h.font = f(True, "FFFFFF", 10); h.fill = fill(col); h.alignment = CENTER
        ws.cell(5, MC0 + k).fill = fill(col)
    ws.row_dimensions[4].height = 14; ws.row_dimensions[5].height = 20

    # master flight programme (rows 6-23): two flight slots per shift, each a code row + time row + comment row
    ws.merge_cells("A6:A23"); ws.merge_cells("B6:B23")
    ws["A6"] = "MASTER"; ws["A6"].font = f(True, "FFFFFF", 10); ws["A6"].fill = fill("1D4ED8"); ws["A6"].alignment = CENTER
    ws["B6"] = "Flight programme\n2 flight slots per shift\nflight code, then time,\nthen comment"; ws["B6"].font = f(True, "1E3A8A", 10); ws["B6"].fill = fill("DBEAFE")
    ws["B6"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    for s, sh in enumerate(SHIFTS):
        bg, fg = SHIFT_STYLE[sh]
        for k in range(2):
            r = CODE_ROW(s, k); rt = r + 1; rc = r + 2
            ws.merge_cells(start_row=r, start_column=3, end_row=rc, end_column=3)
            sc = ws.cell(r, 3, f"{sh} {k + 1}"); sc.fill = fill(bg); sc.font = Font(name="Consolas", bold=True, color=fg, size=9); sc.alignment = CENTER
            for c in range(1, LASTC + 1):
                ws.cell(r, c).border = Border(left=thin, right=thin, top=thin if k == 0 else Side(style="hair", color="CBD5E1"),
                                              bottom=None)
                ws.cell(rt, c).border = Border(left=thin, right=thin)
                ws.cell(rc, c).border = Border(left=thin, right=thin, top=None,
                                               bottom=heavy if k == 1 else Side(style="hair", color="CBD5E1"))
            for c in range(D1, D31 + 1):
                cell = ws.cell(r, c); cell.fill = fill("EFF6FF"); cell.font = f(True, "1D4ED8", 10); cell.alignment = Alignment(horizontal="center", vertical="bottom", shrink_to_fit=True)
                tcell = ws.cell(rt, c); tcell.fill = fill("F8FAFC"); tcell.font = f(False, SLATE700, 9)
                tcell.alignment = Alignment(horizontal="center", vertical="top", shrink_to_fit=True); tcell.number_format = "@"
                ccell = ws.cell(rc, c); ccell.fill = fill("FFFBEB"); ccell.font = f(False, "92400E", 9, italic=True)
                ccell.alignment = Alignment(horizontal="center", vertical="top", shrink_to_fit=True); ccell.number_format = "@"
            ws.row_dimensions[r].height = 18; ws.row_dimensions[rt].height = 15; ws.row_dimensions[rc].height = 15
    ws.merge_cells(start_row=6, start_column=MC0, end_row=23, end_column=LASTC)
    plan = "+".join(f"SUMPRODUCT(SUMIFS(ActSort,ActCodesAll,D{CODE_ROW(s, k)}:AH{CODE_ROW(s, k)}))" for s in range(3) for k in range(2))
    short = "+".join(f'SUMPRODUCT(($D$5:$AH$5<>"")*($D${S0 + 1 + s}:$AH${S0 + 1 + s}>0)*((SUMIFS(ActSort,ActCodesAll,$D${CODE_ROW(s, 0)}:$AH${CODE_ROW(s, 0)})+SUMIFS(ActSort,ActCodesAll,$D${CODE_ROW(s, 1)}:$AH${CODE_ROW(s, 1)}))>$D${S0 + 1 + s}:$AH${S0 + 1 + s}))' for s in range(3))
    ws.cell(3, D1 + len(KPI), "=" + short).number_format = ";;;"
    unc = f"{CL(D1 + len(KPI) - 1)}3"; shr = f"{CL(D1 + len(KPI))}3"
    sm = ws.cell(6, MC0, f'="Planned sorties: "&{plan}&CHAR(10)&"No crew: "&{unc}&" shift(s)"&CHAR(10)&"Short of crew: "&{shr}&" shift(s)"')
    sm.alignment = CENTER; sm.font = f(True, "1E3A8A", 10); sm.fill = fill("EFF6FF")
    ws.conditional_formatting.add(f"{CL(MC0)}6", FormulaRule(formula=[f'{unc}>0'], fill=fill("DC2626"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(f"{CL(MC0)}6", FormulaRule(formula=[f'{shr}>0'], fill=fill("F59E0B"), font=Font(color="FFFFFF", bold=True)))
    dvm = DataValidation(type="list", formula1="FlightCodes", allow_blank=True, showErrorMessage=True, errorStyle="stop",
                         errorTitle="Master flight programme", error="Pick a flight from the dropdown. Flight codes come from the Activities sheet.")
    for s_ in range(3):
        for k_ in range(2):
            dvm.add(f"D{CODE_ROW(s_, k_)}:AH{CODE_ROW(s_, k_)}")
    ws.add_data_validation(dvm)
    ws.conditional_formatting.add("D6:AH23", FormulaRule(formula=['D$5=""'], fill=fill(SLATE400), stopIfTrue=True))
    # a no-flying day copied from the Master Planner: "NO FLY: reason" in the AM 1 comment row greys the day's programme
    ws.conditional_formatting.add("D8:AH8", FormulaRule(formula=['LEFT(D$8,6)="NO FLY"'], fill=fill("B91C1C"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    ws.conditional_formatting.add("D6:AH23", FormulaRule(formula=['LEFT(D$8,6)="NO FLY"'], fill=fill("E5E7EB"), font=Font(color=SLATE400), stopIfTrue=True))
    warn = DifferentialStyle(font=Font(color="FFFFFF", bold=True), fill=fill("DC2626"),
                             numFmt=NumberFormat(numFmtId=200, formatCode='"⚠ "@'))
    low = DifferentialStyle(font=Font(color="FFFFFF", bold=True), fill=fill("F59E0B"),
                            numFmt=NumberFormat(numFmtId=201, formatCode='"! "@'))
    for s in range(3):
        r1, r2, cr = CODE_ROW(s, 0), CODE_ROW(s, 1), S0 + 1 + s
        planned = f"(SUMIFS(ActSort,ActCodesAll,D${r1})+SUMIFS(ActSort,ActCodesAll,D${r2}))"
        for rr in (r1, r2):
            rng = f"D{rr}:AH{rr}"
            ws.conditional_formatting.add(rng, FormulaRule(formula=['D$5=""'], fill=fill(SLATE400), stopIfTrue=True))
            ws.conditional_formatting.add(rng, Rule(type="expression", dxf=warn, stopIfTrue=True, formula=[f'AND(D{rr}<>"",N(D${cr})=0)']))
            ws.conditional_formatting.add(rng, Rule(type="expression", dxf=low, stopIfTrue=True, formula=[f'AND(D{rr}<>"",N(D${cr})<{planned})']))
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'D{rr}<>""'], fill=fill("DCFCE7"), font=Font(color="166534", bold=True), stopIfTrue=True))

    # SDO row: this department's SDO days (from the SDO Plan calendar)
    if HAS_GND:   # ground programme for the office: copied from the Master Planner by the dashboard, or typed in
        ws.merge_cells(start_row=GND_ROW, start_column=1, end_row=GND_ROW, end_column=2)
        lab = ws.cell(GND_ROW, 1, "GROUND PROGRAMME  (from Master Planner)"); lab.font = f(True, "FFFFFF", 9); lab.fill = fill("15803D"); lab.alignment = LEFT
        ws.cell(GND_ROW, 2).fill = fill("15803D")
        sc = ws.cell(GND_ROW, 3, "GND"); sc.font = Font(name="Consolas", bold=True, color="166534", size=9); sc.fill = fill("DCFCE7"); sc.alignment = CENTER
        for c in range(1, LASTC + 1):
            ws.cell(GND_ROW, c).border = BOX
        for d in range(31):
            c = ws.cell(GND_ROW, D1 + d); c.font = f(True, "166534", 8); c.fill = fill("F0FDF4"); c.number_format = "@"
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=False)
        for c in range(MC0, LASTC + 1):
            ws.cell(GND_ROW, c).fill = fill("F0FDF4")
        ws.merge_cells(start_row=GND_ROW, start_column=MC0, end_row=GND_ROW, end_column=LASTC)
        ws.cell(GND_ROW, MC0, "Office ground programme, e.g. CRM Lecture 1400").font = f(False, "166534", 8, italic=True)
        ws.cell(GND_ROW, MC0).alignment = CENTER
        ws.row_dimensions[GND_ROW].height = 30
        grng = f"D{GND_ROW}:AH{GND_ROW}"
        ws.conditional_formatting.add(grng, FormulaRule(formula=['D$5=""'], fill=fill(SLATE400), stopIfTrue=True))
        ws.conditional_formatting.add(grng, FormulaRule(formula=[f'D{GND_ROW}<>""'], fill=fill("DCFCE7"), stopIfTrue=True))
        for dd in seedplan.days():
            if dd.month == m and seedplan.dept_ground(dd, DEPT): ws.cell(GND_ROW, D1 + dd.day - 1, seedplan.dept_ground(dd, DEPT))
    sc0, sr0 = sdo_block(m - 1)
    blk = f"'SDO Plan'!${CL(sc0)}${sr0 + 3}:${CL(sc0 + 6)}${sr0 + 13}"
    bblk = f"'BDO Plan'!${CL(sc0)}${sr0 + 3}:${CL(sc0 + 6)}${sr0 + 13}"
    ws.merge_cells(start_row=SDO_DEPT_ROW, start_column=1, end_row=SDO_DEPT_ROW, end_column=2)
    lab = ws.cell(SDO_DEPT_ROW, 1, "DUTY DAYS  (BDO / SDO, from Duty Roster)" if HAS_DUTY else "SDO DAYS  (set on SDO Plan tab)" if DOES_SDO else f"SDO DAYS  (none, {DEPT} does not do SDO)"); lab.font = f(True, "FFFFFF", 9); lab.fill = fill("BE123C"); lab.alignment = LEFT
    ws.cell(SDO_DEPT_ROW, 2).fill = fill("BE123C")
    sc = ws.cell(SDO_DEPT_ROW, 3, "DUTY" if HAS_DUTY else "SDO"); sc.font = Font(name="Consolas", bold=True, color="881337", size=9); sc.fill = fill("FFE4E6"); sc.alignment = CENTER
    for d in range(31):
        L = CL(D1 + d)
        off = f"{d}+WEEKDAY($A$3,2)-1"
        sfx = f'IF(INDEX({blk},2*INT(({off})/7)+1,MOD({off},7)+1)&""="","","SDO")'
        if HAS_DUTY:
            bfx = f'IF(INDEX({bblk},2*INT(({off})/7)+1,MOD({off},7)+1)&""="","","BDO")'
            c = ws.cell(SDO_DEPT_ROW, D1 + d, f'=IF({L}$5="","",IF(AND({bfx}<>"",{sfx}<>""),"BDO+SDO",{bfx}&{sfx}))')
        else:
            c = ws.cell(SDO_DEPT_ROW, D1 + d, f'=IF({L}$5="","",{sfx})')
        c.font = f(True, "881337", 9); c.fill = fill("FFF1F2"); c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True); c.border = BOX
    for c in range(MC0, LASTC + 1):
        ws.cell(SDO_DEPT_ROW, c).fill = fill("FFF1F2"); ws.cell(SDO_DEPT_ROW, c).border = BOX
    ws.merge_cells(start_row=SDO_DEPT_ROW, start_column=MC0, end_row=SDO_DEPT_ROW, end_column=LASTC)
    ws.cell(SDO_DEPT_ROW, MC0, "Red = nobody planned on that duty" if HAS_DUTY else "Red = nobody planned on SDO" if DOES_SDO else "No SDO duty").font = f(False, "881337", 8, italic=True)
    ws.cell(SDO_DEPT_ROW, MC0).alignment = CENTER
    ws.row_dimensions[SDO_DEPT_ROW].height = 20
    srng = f"D{SDO_DEPT_ROW}:AH{SDO_DEPT_ROW}"
    ws.conditional_formatting.add(srng, FormulaRule(formula=['D$5=""'], fill=fill(SLATE400), stopIfTrue=True))
    redfx = f'AND(D{SDO_DEPT_ROW}<>"",N(D${SDO_ROW})=0)'
    if HAS_DUTY:
        redfx = (f'OR(AND(ISNUMBER(SEARCH("SDO",D{SDO_DEPT_ROW})),COUNTIF(D${G0}:D${GL},"SDO")=0),'
                 f'AND(ISNUMBER(SEARCH("BDO",D{SDO_DEPT_ROW})),COUNTIF(D${G0}:D${GL},"BDO")=0))')
    ws.conditional_formatting.add(srng, FormulaRule(formula=[redfx], fill=fill("DC2626"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(srng, FormulaRule(formula=[f'D{SDO_DEPT_ROW}<>""'], fill=fill("DCFCE7"), font=Font(color="166534", bold=True), stopIfTrue=True))

    # filter header row
    for c in range(1, LASTC + 1):
        ws.cell(FILTER_ROW, c).fill = fill(SLATE900)
    fh = ws.cell(FILTER_ROW, 1, "Callsign"); fh.font = f(True, "FFFFFF", 9); fh.alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(FILTER_ROW, 2, "Empty slots hidden: Data > Reapply shows new crew").font = f(False, SLATE400, 8, italic=True)
    ws.row_dimensions[FILTER_ROW].height = 18

    # crew blocks
    for i in range(N):
        r0 = G0 + PER * i; rr = 6 + i
        band = SANS_FILL if is_sans(i) else "FFFFFF" if i % 2 == 0 else SLATE50
        ws.merge_cells(start_row=r0, start_column=2, end_row=r0 + 1, end_column=2)
        for s in range(PER):
            a = ws.cell(r0 + s, 1, f'=IF(Roster!$C${rr}="","",IF(Roster!$B${rr}="","#"&{i + 1},Roster!$B${rr}&""))')
            a.font = f(True, SLATE700, 9); a.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
            if s:
                a.number_format = ";;;"   # same value on all four rows so the filter keeps the block together
        b = ws.cell(r0, 2, f'=IF(Roster!$C${rr}="","",Roster!$C${rr})'); b.font = f(True, SLATE800, 11)
        b.alignment = Alignment(horizontal="left", vertical="bottom", indent=1, wrap_text=True)
        rl = ws.cell(r0 + 2, 2, f'=IF(Roster!$C${rr}="","",Roster!$F${rr}&IF(OR(Roster!$G${rr}="SANS",Roster!$G${rr}="NSMEN"),"  |  "&Roster!$G${rr},""))'); rl.font = f(False, SLATE500, 9)
        rl.alignment = Alignment(horizontal="left", vertical="top", indent=1, shrink_to_fit=True)
        for s, sh in enumerate(ROW_LABELS):
            r = r0 + s
            bg, fg = SHIFT_STYLE[sh]
            sc = ws.cell(r, 3, sh); sc.fill = fill(bg); sc.font = Font(name="Consolas", bold=True, color=fg, size=9); sc.alignment = CENTER
            bottom = heavy if s == PER - 1 else Side(style="thin", color=SLATE200)
            for c in range(1, LASTC + 1):
                cell = ws.cell(r, c)
                cell.border = Border(left=thin, right=thin, top=Side(style="thin", color=SLATE200) if s else Side(style="thick", color="7E22CE") if i == REG_N else heavy, bottom=bottom)
                if c != 3 and (c <= 2 or not is_sans(i) or c > D31):
                    cell.fill = fill(WD_TINT if s == PER - 1 and D1 <= c <= D31 else band)
                if D1 <= c <= D31:
                    cell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True); cell.font = f(False, SLATE700, 10)
            ws.row_dimensions[r].height = 19
        blk3 = f"$D{r0}:$AH{r0 + 2}"   # the three shift rows: flights, hours and duty never sit in WD
        for k, (lab, cat, mode, col) in enumerate(METRICS):
            hc = f"{CL(MC0 + k)}$3"
            if mode == "entries":
                fx = f"SUMPRODUCT(SUMIFS(ActSort,ActCodesAll,{blk3},ActCat,{hc}))"
            elif mode == "hours":
                fx = f"SUMPRODUCT(SUMIFS(ActHrs,ActCodesAll,{blk3},ActCat,{hc}))"
            elif mode == "duty":
                rows = "+".join(f"($D{r0 + s}:$AH{r0 + s}=\"{dc}\")" for s in range(3) for dc in DUTY_CODES)
                fx = f"SUMPRODUCT(--(({rows})>0))"
            else:
                rows = "+".join(f"COUNTIFS(ActCodesAll,$D{r0 + s}:$AH{r0 + s},ActCat,{hc})" for s in range(PER))
                fx = f"SUMPRODUCT(--(({rows})>0))"
            ws.merge_cells(start_row=r0, start_column=MC0 + k, end_row=r0 + PER - 1, end_column=MC0 + k)
            c = ws.cell(r0, MC0 + k, f'=IF($A{r0}="","",{fx})')
            c.alignment = CENTER; c.font = f(True, SLATE900, 12)
            c.number_format = "0.0;-0.0;-" if mode == "hours" else "0;-0;-"
        if i not in crew:
            for s in range(PER):
                ws.row_dimensions[r0 + s].hidden = not NOFILTER
    nonblank_filter(ws, f"A{FILTER_ROW}:A{GL}", 0)

    # daily summary (all departments)
    ws.cell(S0, 2, "DAILY SUMMARY").font = f(True, SLATE900, 11)
    def persons(L, cats):
        parts = []
        for s in range(PER):
            rng = f"{L}${G0 + s}:{L}${GL - (PER - 1) + s}"
            parts += [f"COUNTIFS(ActCodesAll,{rng},ActCat,{catref(c)})" for c in cats]
        return f"SUMPRODUCT(($C${G0}:$C${GL - (PER - 1)}=\"AM\")*(({'+'.join(parts)})>0))"
    def persons_duty(L):
        parts = [f'({L}${G0 + s}:{L}${GL - (PER - 1) + s}="{dc}")' for s in range(3) for dc in DUTY_CODES]
        return f"SUMPRODUCT(($C${G0}:$C${GL - (PER - 1)}=\"AM\")*(({'+'.join(parts)})>0))"
    for lab in SUMMARY:
        r = SUM_ROW[lab]
        ws.cell(r, 2, lab).font = f(lab in ("Total flights", "SDO on duty", "Available crew"), SLATE800, 10)
        ws.cell(r, 2).alignment = LEFT
        for c in (2, 3):
            ws.cell(r, c).border = BOX; ws.cell(r, c).fill = fill(SLATE50)
        for d in range(31):
            L = CL(D1 + d)
            if lab.startswith("Flights "):
                sh = lab.split()[1]
                fx = f'SUMPRODUCT(($C${G0}:$C${GL}="{sh}")*SUMIFS(ActSort,ActCodesAll,{L}${G0}:{L}${GL},ActCat,{catref("Flight")}))'
            elif lab == "Total flights":
                fx = f"SUM({L}{S0 + 1}:{L}{S0 + 3})"
            elif lab == "Fly hrs":
                fx = f'SUMPRODUCT(SUMIFS(ActHrs,ActCodesAll,{L}${G0}:{L}${GL},ActCat,{catref("Flight")}))'
            elif lab == "SDO on duty":
                fx = persons_duty(L)
            elif lab == "Gnd duty (crew)":
                fx = persons(L, ["Ground Duty"])
            elif lab == "On leave (crew)":
                fx = persons(L, ["Leave"])
            else:
                fx = '(COUNTIFS(CrewName,"<>")-COUNTIFS(CrewName,"<>",CrewStatus,"Inactive"))-' + persons(L, ["Leave", "Off", "Medical", "Oversea Duty"])
            c = ws.cell(r, D1 + d, f'=IF({L}$5="","",{fx})')
            c.alignment = CENTER; c.border = BOX
            c.font = f(lab in ("Total flights", "SDO on duty", "Available crew"), SLATE700, 9)
            c.number_format = "0.0;-0.0;-" if lab == "Fly hrs" else "0;-0;-"
        ws.row_dimensions[r].height = 16
    ws.conditional_formatting.add(f"D{SDO_ROW}:AH{SDO_ROW}", FormulaRule(
        formula=[f'AND(D${SDO_DEPT_ROW}<>"",D{SDO_ROW}=0)'], fill=fill("F43F5E"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(f"D{SDO_ROW}:AH{SDO_ROW}", FormulaRule(
        formula=[f'AND(D${SDO_DEPT_ROW}="",N(D{SDO_ROW})>0)'], fill=fill("F59E0B"), font=Font(color="FFFFFF", bold=True)))

    grid = f"D{G0}:AH{GL}"
    # a crew cell with a comment on the Cell Notes sheet gets a purple outline (added first so the activity colours still apply)
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'COUNTIF(NoteKey,D$5&"|"&ROW())>0'], border=Border(left=NOTE_SIDE, right=NOTE_SIDE, top=NOTE_SIDE, bottom=NOTE_SIDE), font=Font(underline="single")))
    apply_cat_cf(ws, grid, f"D{G0}")
    ws.conditional_formatting.add(f"D{G0}:AH{S0 + len(SUMMARY)}", FormulaRule(formula=['D$5=""'], fill=fill(SLATE400), stopIfTrue=True))
    ws.conditional_formatting.add(grid, FormulaRule(formula=['COUNTIF(Holidays,D$5)>0'], fill=fill("FFE4E6"), stopIfTrue=True))
    ws.conditional_formatting.add(grid, FormulaRule(formula=['WEEKDAY(D$5,2)>5'], fill=fill("E2E8F0"), stopIfTrue=True))
    ws.conditional_formatting.add("D4:AH5", FormulaRule(formula=['AND(D$5<>"",COUNTIF(Holidays,D$5)>0)'], fill=fill("9F1239"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add("D4:AH5", FormulaRule(formula=['AND(D$5<>"",WEEKDAY(D$5,2)>5)'], fill=fill(SLATE700), font=Font(color=AMBER300, bold=True)))
    ws.conditional_formatting.add(f"B{G0}:B{GL}", FormulaRule(
        formula=[f'INDEX(CrewStatus,INT((ROW()-{G0})/{PER})+1)="Inactive"'], font=Font(color=SLATE400, italic=True, strike=True), stopIfTrue=True))
    cat_cf(ws, f"B{G0}:B{GL}", f"INT((ROW()-{G0})/{PER})+1")
    dv = DataValidation(type="list", formula1=f'IF($C{G0}="WD",WdCodes,ActCodes)', allow_blank=True, showErrorMessage=True, errorStyle="stop",
                        errorTitle="Unknown activity", error="Pick an activity from the dropdown. New activities go on the Activities sheet. The WD row takes ground and admin items only.")
    dv.add(grid); ws.add_data_validation(dv)

    ws.column_dimensions["A"].width = 11; ws.column_dimensions["B"].width = 27; ws.column_dimensions["C"].width = 8
    for c in range(D1, D31 + 1):
        ws.column_dimensions[CL(c)].width = 10.5
    for k in range(len(METRICS)):
        ws.column_dimensions[CL(MC0 + k)].width = 9.5
    ws.freeze_panes = "D6" if SMALLFREEZE else f"D{G0}"
    ws.sheet_view.zoomScale = 100
    if COMPACT:
        # tighter month tab for smaller / high-scaling screens: lower rows, narrower days, 80% zoom
        ws["A1"].font = f(True, "FFFFFF", 12); ws.row_dimensions[1].height = 20; ws.row_dimensions[2].hidden = True
        ws.row_dimensions[4].height = 12; ws.row_dimensions[5].height = 16
        for s_ in range(3):
            for k in range(2):
                r = CODE_ROW(s_, k); ws.row_dimensions[r].height = 14; ws.row_dimensions[r + 1].height = 12; ws.row_dimensions[r + 2].height = 12
        ws["B6"] = "Flight programme: code, time, comment"; ws["B6"].font = f(True, "1E3A8A", 9)
        ws.row_dimensions[SDO_DEPT_ROW].height = 15; ws.row_dimensions[FILTER_ROW].height = 14
        for r in range(G0, GL + 1):
            if ws.row_dimensions[r].height == 19: ws.row_dimensions[r].height = 16
        for r in range(S0, S0 + len(SUMMARY) + 1):
            if ws.row_dimensions[r].height == 16: ws.row_dimensions[r].height = 14
        ws.column_dimensions["B"].width = 22
        for c in range(D1, D31 + 1):
            ws.column_dimensions[CL(c)].width = 9
        ws.sheet_view.zoomScale = 80
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:5"

    if m in (10, 11):   # sample plan (seedplan.py): crew cells, and the MASTER rows as the Master Planner gives them to this department
        ids = {c[0]: k for k, c in crew.items()}
        for (pid, dd, sh), code in seedplan.plan(DEPT).items():
            if pid in ids and dd.month == m:
                ws.cell(G0 + PER * ids[pid] + SHIFTS.index(sh), D1 + dd.day - 1, code)
        for dd in seedplan.days():
            if dd.month != m: continue
            for rr, (code, tm, nt) in enumerate(seedplan.dept_slots(dd, DEPT)):
                s_, k = divmod(rr, 2)
                if code: ws.cell(CODE_ROW(s_, k), D1 + dd.day - 1, code)
                if tm: ws.cell(CODE_ROW(s_, k) + 1, D1 + dd.day - 1, tm)
                if nt: ws.cell(CODE_ROW(s_, k) + 2, D1 + dd.day - 1, nt)

# sample duty days (from seedplan.py, the same as the Duty Roster): this department's BDO / SDO days in October and November
for code, ws_ in (("SDO", sdo), ("BDO", bdo)):
    if ws_ is None: continue
    for dd, dep in seedplan.DUTY[code].items():
        if dep != DEPT: continue
        c0, r0 = sdo_block(dd.month - 1); n = dd.day - 1 + dt.date(YEAR, dd.month, 1).weekday()
        ws_.cell(r0 + 3 + 2 * (n // 7), c0 + n % 7, dep)

# ---------------------------------------------------------------- Dashboard
LASTD = 14
banner(dash, f"LODESTAR  |  {DEPT.upper()}  |  FLIGHT & OPS PLANNING DASHBOARD",
       "Pick a year and month to see its totals. Planning is done on the month tabs: click Open planner.", LASTD)
dash.row_dimensions[3].height = 8
dash["A4"] = "YEAR"; dash["A4"].font = f(True, SLATE500, 9)
dash["B4"] = "MONTH"; dash["B4"].font = f(True, SLATE500, 9)
dash["A5"] = YEAR; dash["A5"].font = f(True, "1D4ED8", 14); dash["A5"].fill = INPUT_FILL; dash["A5"].border = BOX; dash["A5"].alignment = CENTER
dash["B5"] = MONTHS_FULL[9]; dash["B5"].font = f(True, "1D4ED8", 14); dash["B5"].fill = INPUT_FILL; dash["B5"].border = BOX; dash["B5"].alignment = LEFT
if not NOCOMMENT: dash["A5"].comment = Comment("Planning year for the whole workbook. For a new year, save a copy of the file, change this, and clear the month grids and SDO Plan.", "Claude")
dash.merge_cells("C5:D5")
dash["C5"] = f'=HYPERLINK("#\'"&LEFT($B$5,3)&"\'!D{G0}","Open "&$B$5&" planner  >")'
dash["C5"].font = Font(name=FONT, bold=True, color="FFFFFF", size=11, underline="single"); dash["C5"].fill = fill("1D4ED8"); dash["C5"].alignment = CENTER
dash.merge_cells("E5:F5")
dash["E5"] = '=HYPERLINK("#\'SDO Plan\'!A1","SDO Plan  >")'
dash["E5"].font = Font(name=FONT, bold=True, color="FFFFFF", size=11, underline="single"); dash["E5"].fill = fill("BE123C"); dash["E5"].alignment = CENTER
dash.row_dimensions[5].height = 24
dash["G4"] = "DEPARTMENT (this file)"; dash["G4"].font = f(True, SLATE500, 9)
dash.merge_cells("G5:I5")
dash["G5"] = DEPT; dash["G5"].font = f(True, "BE123C", 14); dash["G5"].fill = INPUT_FILL; dash["G5"].border = BOX; dash["G5"].alignment = LEFT
if not NOCOMMENT: dash["G5"].comment = Comment("The department that owns this file. The dashboard uses it to label this file when several departments are open.", "Claude")
name("ThisDept", "Dashboard!$G$5")
name("PlanYear", "Dashboard!$A$5")
name("MonthIdx", "Dashboard!$Z$1")
dash["Z1"] = "=MATCH($B$5,{" + ",".join(f'"{x}"' for x in MONTHS_FULL) + "},0)"
dash.column_dimensions["Z"].hidden = True
# layout stamp: the dashboard compares it with the layout it was written for (bump it when cell positions change)
dash["Z2"] = "LODESTAR dept layout 3"; dash.column_dimensions["Z"].hidden = True
dv = DataValidation(type="list", formula1='"' + ",".join(str(y) for y in range(2025, 2036)) + '"'); dv.add("A5"); dash.add_data_validation(dv)
dv = DataValidation(type="list", formula1='"' + ",".join(MONTHS_FULL) + '"'); dv.add("B5"); dash.add_data_validation(dv)
def pick(cellref):
    return "CHOOSE(MonthIdx," + ",".join(f"'{mon}'!{cellref}" for mon in MONTHS) + ")"
kpi_cells = [f"{CL(D1 + k)}3" for k in range(len(KPI))]
dash.row_dimensions[6].height = 8
positions = [(1, 1), (2, 1), (3, 1), (4, 2), (6, 2), (8, 2), (11, 2)]
for k, (lab, fx, accent, nf) in enumerate(KPI):
    c0, w = positions[k]
    formula = '=(COUNTIFS(CrewName,"<>")-COUNTIFS(CrewName,"<>",CrewStatus,"Inactive"))' if lab == "Active Crew" else f"={pick(kpi_cells[k])}"
    card(dash, 7, c0, w, lab, formula, accent, nf)

T0 = 10
dash.cell(T0, 1, '="CREW TOTALS  |  "&UPPER($B$5)&" "&$A$5').font = f(True, SLATE900, 11)
dash.cell(T0, 4, "Empty slots are hidden: Data > Reapply shows new crew").font = f(italic=True, color=SLATE500, size=8)
header(dash, T0 + 1, 1, ["Callsign", "Crew Member", "Role", "Status"] + [m[0] for m in METRICS])
for k, (lab, cat, mode, col) in enumerate(METRICS):
    dash.cell(T0 + 1, 5 + k).fill = fill(col)
for i in range(N):
    r = T0 + 2 + i; rr = 6 + i; gr = G0 + PER * i
    dash.cell(r, 1, f'=IF(Roster!$B${rr}="","",Roster!$B${rr})').font = f(True, SLATE900, 10)
    dash.cell(r, 2, f'=IF(Roster!$C${rr}="","",Roster!$C${rr})').font = f(True, SLATE800, 10)
    dash.cell(r, 3, f'=IF(Roster!$C${rr}="","",Roster!$F${rr}&"")').font = f(False, SLATE700, 10)
    dash.cell(r, 4, f'=IF(Roster!$C${rr}="","",Roster!$G${rr}&"")').font = f(False, SLATE500, 10)
    for k, (lab, cat, mode, col) in enumerate(METRICS):
        c = dash.cell(r, 5 + k, f'=IF($B{r}="","",{pick(CL(MC0 + k) + str(gr))})')
        c.alignment = CENTER; c.number_format = "0.0;-0.0;-" if mode == "hours" else "0;-0;-"; c.font = f(True, SLATE900, 11)
    for c in range(1, 5 + len(METRICS)):
        dash.cell(r, c).border = BOX
        if i % 2 or is_sans(i): dash.cell(r, c).fill = fill(SANS_FILL if is_sans(i) else SLATE50)
    if i not in crew:
        dash.row_dimensions[r].hidden = not NOFILTER
TL = T0 + 1 + N
nonblank_filter(dash, f"A{T0 + 1}:I{TL}", 1)
tr = TL + 1
dash.cell(tr, 2, "Squadron total").font = f(True)
for k, (lab, cat, mode, col) in enumerate(METRICS):
    c = dash.cell(tr, 5 + k, f"=SUM({CL(5 + k)}{T0 + 2}:{CL(5 + k)}{TL})"); c.font = f(True); c.alignment = CENTER
    c.number_format = "0.0" if mode == "hours" else "0"
for c in range(1, 5 + len(METRICS)):
    dash.cell(tr, c).border = Border(top=Side(style="medium", color=SLATE900)); dash.cell(tr, c).fill = fill(SLATE100)
dash.conditional_formatting.add(f"A{T0 + 2}:D{TL}", FormulaRule(formula=['$D' + str(T0 + 2) + '="Inactive"'], font=Font(color=SLATE400, italic=True, strike=True), stopIfTrue=True))
cat_cf(dash, f"B{T0 + 2}:B{TL}", f"ROW()-{T0 + 1}")

# activity legend on the right of the department table
LG = 11
dash.cell(T0, LG, "ACTIVITY LEGEND").font = f(True, SLATE900, 11)
header(dash, T0 + 1, LG, ["Code", "Description", "", "Hrs"])
dash.merge_cells(start_row=T0 + 1, start_column=LG + 1, end_row=T0 + 1, end_column=LG + 2)
for i in range(20):
    r = T0 + 2 + i; ar = 6 + i
    dash.cell(r, LG, f'=IF(Activities!$A${ar}="","",Activities!$A${ar})').alignment = CENTER
    dash.merge_cells(start_row=r, start_column=LG + 1, end_row=r, end_column=LG + 2)
    dash.cell(r, LG + 1, f'=IF(Activities!$A${ar}="","",Activities!$B${ar}&"")').font = f(False, SLATE700, 9)
    h = dash.cell(r, LG + 3, f'=IF(Activities!$A${ar}="","",Activities!$D${ar})'); h.number_format = "0.0;-0.0;"; h.font = f(False, SLATE700, 9)
    h.alignment = CENTER
    dash.cell(r, LG).font = f(True, SLATE900, 9)
apply_cat_cf(dash, f"K{T0 + 2}:K{T0 + 21}", f"K{T0 + 2}")

# year overview
Y0 = tr + 3
dash.cell(Y0, 1, '="YEAR OVERVIEW  |  "&$A$5').font = f(True, SLATE900, 11)
dash.merge_cells(start_row=Y0, start_column=1, end_row=Y0, end_column=3)
ylabels = [("Flights", 0, "0"), ("Flying Hrs", 1, "0.0"), ("Ground Duties", 2, "0"), ("SDO Planned", 3, "0"), ("SDO Days Uncovered", 5, "0"), ("Uncrewed Shifts", 6, "0")]
header(dash, Y0 + 1, 1, ["Month"] + [y[0] for y in ylabels])
for mi, mon in enumerate(MONTHS):
    r = Y0 + 2 + mi
    lk = dash.cell(r, 1, f'=HYPERLINK("#\'{mon}\'!D{G0}","{MONTHS_FULL[mi]}")'); lk.font = Font(name=FONT, color="1D4ED8", underline="single", size=10)
    for j, (lab, k, nf) in enumerate(ylabels):
        c = dash.cell(r, 2 + j, f"='{mon}'!{kpi_cells[k]}"); c.alignment = CENTER; c.font = f(False, "15803D", 10)
        c.number_format = "0.0;-0.0;-" if "." in nf else "0;-0;-"
    for c in range(1, 2 + len(ylabels)):
        dash.cell(r, c).border = BOX
yr = Y0 + 14
dash.cell(yr, 1, "Year").font = f(True)
for j, (lab, k, nf) in enumerate(ylabels):
    L = CL(2 + j)
    c = dash.cell(yr, 2 + j, f"=SUM({L}{Y0 + 2}:{L}{Y0 + 13})"); c.font = f(True); c.alignment = CENTER; c.number_format = nf
for c in range(1, 2 + len(ylabels)):
    dash.cell(yr, c).border = Border(top=Side(style="medium", color=SLATE900)); dash.cell(yr, c).fill = fill(SLATE100)
dash.conditional_formatting.add(f"A{Y0 + 2}:G{Y0 + 13}", FormulaRule(formula=[f'MonthIdx=ROW()-{Y0 + 1}'], fill=fill("DBEAFE")))
for j, (title, col) in enumerate((("Flights per month", 2), ("Flying hours per month", 3))):
    ch = BarChart(); ch.type = "col"; ch.title = title; ch.style = 10 if j == 0 else 12
    ch.add_data(Reference(dash, min_col=col, min_row=Y0 + 1, max_row=Y0 + 13), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=Y0 + 2, max_row=Y0 + 13))
    ch.height = 6.5; ch.width = 12; ch.legend = None
    if not NOCHART: dash.add_chart(ch, f"I{Y0 + 1 + 14 * j}")
for col, w in zip("ABCDEFGHIJKLMN", [15, 24, 13, 11, 10, 10, 10, 10, 10, 3, 11, 15, 15, 8]):
    dash.column_dimensions[col].width = w
dash.freeze_panes = "A7"

# ---------------------------------------------------------------- Daily view
banner(daily, f"DAILY OPS VIEW  |  {DEPT.upper()}", "Type a date to see who is on what. Shows today by default.", 8)
daily["A4"] = "DATE"; daily["A4"].font = f(True, SLATE500, 9)
daily["A5"] = "=TODAY()"; daily["A5"].number_format = "ddd dd-mmm-yyyy"
daily["A5"].font = f(True, "1D4ED8", 12); daily["A5"].fill = INPUT_FILL; daily["A5"].border = BOX
daily.merge_cells("A5:B5")
daily["C5"] = '=IF(NOT(' + INYR("A5") + '),"Date is outside planning year "&PlanYear&" and January "&(PlanYear+1),"")'; daily["C5"].font = f(True, "DC2626")
dsum = [("Flights", "Total flights", "1D4ED8", "0"), ("Flying Hrs", "Fly hrs", "0F766E", "0.0"),
        ("SDO on Duty", "SDO on duty", "BE123C", "0"), ("On Leave", "On leave (crew)", SLATE700, "0"),
        ("Available Crew", "Available crew", "15803D", "0")]
for k, (lab, key, accent, nf) in enumerate(dsum):
    sr = SUM_ROW[key]
    ref = "CHOOSE(" + MIX("$A$5") + "," + ",".join(f"'{mon}'!$D${sr}:$AH${sr}" for mon in MONTHS13) + ")"
    card(daily, 7, [1, 3, 4, 5, 6][k], 2 if k == 0 else 1, lab,
         f'=IF(NOT('+INYR('$A$5')+'),"",INDEX({ref},1,DAY($A$5)))', accent, nf)
sdoref = "CHOOSE(" + MIX("$A$5") + "," + ",".join(f"'{mon}'!$D${SDO_DEPT_ROW}:$AH${SDO_DEPT_ROW}" for mon in MONTHS13) + ")"
if HAS_DUTY:
    card(daily, 7, 7, 1, "Duty Day", f'=IF(NOT('+INYR('$A$5')+'),"",IF(INDEX({sdoref},1,DAY($A$5))="","No",INDEX({sdoref},1,DAY($A$5))))', "BE123C", "@")
else:
    card(daily, 7, 7, 1, "SDO Day", f'=IF(NOT('+INYR('$A$5')+'),"",IF(INDEX({sdoref},1,DAY($A$5))="","No","Yes"))', "BE123C", "@")
header(daily, 10, 1, ["Callsign", "Crew Member", "Role", "AM", "PM", "NIGHT", "WD"])
chooser = "CHOOSE(" + MIX("$A$5") + "," + ",".join(f"'{mon}'!$D${G0}:$AH${GL}" for mon in MONTHS13) + ")"
for i in range(N):
    r = 11 + i; rr = 6 + i
    daily.cell(r, 1, f'=IF(Roster!$B${rr}="","",Roster!$B${rr})').font = f(True, SLATE900, 10)
    daily.cell(r, 2, f'=IF(Roster!$C${rr}="","",Roster!$C${rr})').font = f(True, SLATE800, 10)
    daily.cell(r, 3, f'=IF(Roster!$C${rr}="","",Roster!$F${rr}&IF(OR(Roster!$G${rr}="SANS",Roster!$G${rr}="NSMEN")," | "&Roster!$G${rr},""))').font = f(False, SLATE700, 10)
    for s in range(PER):
        c = daily.cell(r, 4 + s, f'=IF(OR($B{r}="",NOT('+INYR('$A$5')+f')),"",INDEX({chooser},{PER * i + s + 1},DAY($A$5))&"")')
        c.alignment = CENTER; c.font = f(False, SLATE700, 10)
    for c in range(1, 8):
        daily.cell(r, c).border = BOX
        if is_sans(i) and c <= 3: daily.cell(r, c).fill = fill(SANS_FILL)
    if i not in crew:
        daily.row_dimensions[r].hidden = not NOFILTER
nonblank_filter(daily, f"A10:G{10 + N}", 1)
apply_cat_cf(daily, f"D11:G{10 + N}", "D11")
daily.conditional_formatting.add(f"A11:B{10 + N}", FormulaRule(formula=['INDEX(CrewStatus,ROW()-10)="Inactive"'], font=Font(color=SLATE400, italic=True, strike=True), stopIfTrue=True))
cat_cf(daily, f"B11:B{10 + N}", "ROW()-10")
for col, w in zip("ABCDEFGH", [12, 26, 14, 15, 15, 15, 15, 3]):
    daily.column_dimensions[col].width = w
daily.freeze_panes = "A11"

# ---------------------------------------------------------------- Guide
banner(guide, "LODESTAR FLIGHT & OPS PLANNER: HOW IT WORKS", "Offline, macro-free. Desktop Excel 2010 or later.", 1)
lines = [
    ("ONE FILE PER DEPARTMENT", "h"),
    ("Each department keeps its own copy of this workbook, with its own roster, flight programme and SDO days. The department name is on the Dashboard (cell G5).", None),
    ("To see other departments, open the planner dashboard (the .html file), open your own Excel, then click 'Add other departments' and pick their Excel files. Each one shows as a read-only tab, and the SDO Plan tab combines everyone's SDO days.", None),
    *([("Duty days (BDO = Base Duty Officer, SDO = Squadron Duty Orderly) come from the Duty Planner's Duty Roster file. In the dashboard, open your own file, Add other departments and pick the Duty Roster, then on the Duty Planner tab click 'Copy my duty days from the Duty Roster' and Save to Excel.", None)] if HAS_DUTY else [
    ("SDO days come from the SDO Planner's SDO Roster file. In the dashboard, open your own file, Add other departments and pick the SDO Roster, then on the SDO Plan tab click 'Copy my SDO days from the SDO Roster' and Save to Excel.", None),
    ]),
    ("New department: copy an empty file, change the Department name on the Dashboard, and fill in its roster.", None),
    ("", None),
    ("SHEETS", "h"),
    ("Dashboard: pick a year and month to see KPI cards, crew totals, the activity legend and the year overview. 'Open planner' jumps to that month.", None),
    ("Jan to Dec: the planning matrix. Each crew member has AM, PM and NIGHT rows. Click a day cell and pick an activity from the dropdown.", None),
    *([("BDO Plan and SDO Plan: year calendars of your department's BDO and SDO days, copied from the Duty Roster. Each month tab shows them in the pink DUTY DAYS row and the table on the right lists the dates.", None),
       ("DUTY DAYS row (month tabs): green when someone in your department has that duty planned (activity BDO or SDO), red when nobody does.", None)] if HAS_DUTY else [
    ("SDO Plan: a year calendar. On each day your department has SDO duty, pick your department in the dropdown. Each month tab marks those days in the pink SDO row and the table on the right lists the dates.", None),
    ("SDO row (month tabs): green when someone in your department has SDO planned that day, red when nobody does. In the daily summary, SDO on duty turns amber when someone is on SDO on a day that isn't yours.", None),
    ]),
    ("Master rows (top of each month tab): the flight programme. Each shift has two flight slots (AM 1, AM 2, PM 1 ...); pick a flight from the dropdown in each. Leave slot 2 empty if only one flight is planned.", None),
    ("Flight timing: type it in the grey row under each master flight, in any form (0830, 0830-1030, TBC). Display only.", None),
    ("Flight comments: type any short note in the cream row under the time (e.g. 'Wx dependent', 'Need 2nd crew'). Display only; long notes shrink to fit the cell.", None),
    ("More than two flights in a shift, or one person flying twice in a shift: use a combined code such as 2xFLT-1.5 (Activities sheet, Sorties = 2). It counts as 2 flights and its total hours.", None),
    ("Master warnings: red with a warning sign = flights planned but nobody in that shift is flying. Amber with ! = fewer crew sorties than planned sorties. Green = covered.", None),
    ("Daily View: type any date to see every crew member's AM, PM and NIGHT activity, that day's totals and whether it's an SDO day.", None),
    ("Roster: the crew list (callsign, name with rank, initials, Cat Status, role, status). Activities: the dropdown catalog, categories and public holidays.", None),
    ("", None),
    ("ADDING AND REMOVING CREW", "h"),
    ("Add: type them in the next empty Roster row. They appear on every month tab, the Dashboard and Daily View.", None),
    ("Regular crew go in rows 6-50 (Status Active). SANS and NSmen go in the purple rows 51-65 (Status SANS or NSMEN), so they always show at the bottom of every planner.", None),
    ("Cat Status colours the crew member's name: CAT A black, CAT B green, CAT C yellow, CAT D red, TRAINEE magenta.", None),
    ("Empty crew slots are hidden by a filter. After adding someone in Excel, click Data > Reapply (or Ctrl+Alt+L) on each month tab, the Dashboard and Daily View to show them. Crew added through the dashboard show up straight away. There are 60 slots.", None),
    ("Remove: set Status to Inactive. Their name is struck through and the crew count drops, but their past plans and totals are kept.", None),
    ("Don't delete or re-order Roster rows. Month tabs link to the Roster by row, so deleting a row moves everyone's plan onto the wrong person.", "warn"),
    ("", None),
    ("ADDING ACTIVITIES", "h"),
    ("Type a new code on the Activities sheet directly under the last one (no gaps), pick its category and enter its flight hours. It appears in every dropdown.", None),
    ("", None),
    ("HOW TOTALS ARE COUNTED", "h"),
    ("Flights: one per Flight-category cell, so flying AM and NIGHT on the same day is 2 flights. Fly Hrs adds each activity's Flight Hrs.", None),
    ("Gnd Duty, SDO and Leave are counted in days. A day counts once if any of AM, PM or NIGHT has that category, so a 24h SDO filled into all three rows is 1 SDO.", None),
    ("", None),
    ("NEW YEAR" if REAL else "SAMPLE DATA", "h"),
    *([] if REAL else [("The crew, October entries and October SDO days are samples. Clear them (select the cells and press Delete) before real use.", None)]),
    ("New year: in the dashboard Ops Room, Master Planner menu > Start new year. It archives this year's files and moves every file to the next year (crew, activities, currency dates and open down-chits carry over).", None),
]
for i, (t, kind) in enumerate(lines):
    c = guide.cell(4 + i, 1, t)
    c.font = f(True, SLATE900, 12) if kind == "h" else f(True, "DC2626", 11) if kind == "warn" else f(False, SLATE700, 11)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
guide.column_dimensions["A"].width = 140

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.move_sheet("Dashboard", offset=-1)
wb.calculation.fullCalcOnLoad = True
wb.active = wb.sheetnames.index("Dashboard")
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == "Dashboard"
wb.save(OUT)
if WEBCLEAN:
    # Excel Online test: drop the empty cached values openpyxl writes on every formula cell, and the empty workbook-protection tag
    import zipfile, shutil
    tmp = OUT + ".tmp"
    with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.endswith(".xml"):
                data = data.replace(b"<v></v>", b"").replace(b"<v />", b"").replace(b"<workbookProtection/>", b"").replace(b"<workbookProtection />", b"")
            zo.writestr(it, data)
    shutil.move(tmp, OUT)
print("saved", OUT, wb.sheetnames)
