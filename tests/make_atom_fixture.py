#!/usr/bin/env python3
"""Synthetic ATOM export for the tests (made-up people; never copy a real export into this repo).
python3 tests/make_atom_fixture.py OUT.xlsx [--created "2026-10-25 12:41:33"] [--variant moved|partial]
Structure copied from the real exports: row 1 group titles (merged), row 2 column names, row 3 a status band, people from row 4,
band rows (' -> NSmen') between groups. Dates are written as text 'dd Mon yy'; one is a real Excel date."""
import sys, datetime as dt, openpyxl
from openpyxl.styles import Alignment

out = sys.argv[1]
created = dt.datetime(2026, 10, 25, 12, 41, 33)
variant = ""
a = sys.argv[2:]
if "--created" in a: created = dt.datetime.strptime(a[a.index("--created") + 1], "%Y-%m-%d %H:%M:%S")
if "--variant" in a: variant = a[a.index("--variant") + 1]

# group -> columns
PILOT_COLS = [("Flight Currencies", ["Day Capt", "Day Curr", "IRT", "Instructor Check", "Night Capt", "Night Curr"]), ("SIM Currencies", ["SIM"]),
              ("LSS Currencies", ["ACE", "APT", "SVT"]), ("Mission Currencies", ["MAS", "SAL"]),
              ("SIM Package (Oct 25 - Oct 26)", ["SIM Package"]), ("Flying Hours Clocked", ["Current Month"]), ("SATR Details", ["Level", "Remarks"])]
OTHER_COLS = [("Flight Currencies", ["Day Curr", "Night Curr"]), ("SIM Currencies", ["SIM"]), ("LSS Currencies", ["ACE", "APT"]), ("Mission Currencies", ["MAS", "SAL"]),
              ("SIM Package (Oct 25 - Oct 26)", ["SIM Package"]), ("Flying Hours Clocked", ["Current Month"]), ("SATR Details", ["Level", "Remarks"])]
if variant == "moved":   # an added column first in a group, a reordered group
    PILOT_COLS = [("Flight Currencies", ["NVG Curr"] + PILOT_COLS[0][1]), PILOT_COLS[3], PILOT_COLS[1], PILOT_COLS[2]] + PILOT_COLS[4:]

D = lambda s: s   # text date such as '16 Nov 26'
XL = dt.datetime(2027, 1, 22)   # stored as a real Excel date
OK, DUE, LAP = "22 Jan 27", "05 Nov 26", "16 Oct 26"
# per person: {column: value}; anything missing = '-' (not held). Day of reference in the tests: 2026-10-26.
PEOPLE = {
    "PILOT": [("Active", "A NGO", "A", {"Day Capt": LAP, "Day Curr": "16 Nov 26", "IRT": "-", "Instructor Check": "N/A", "Night Capt": LAP, "Night Curr": "27 Oct 26", "SIM": OK, "ACE": " ", "APT": "soon", "SVT": XL, "MAS": "30 Jun 31", "SAL": "-", "NVG Curr": OK}),
              ("Active", "HO KM", "B", {"Day Capt": OK, "Day Curr": OK, "IRT": OK, "Instructor Check": OK, "Night Capt": OK, "Night Curr": OK, "SIM": OK, "ACE": OK, "APT": OK, "SVT": OK, "MAS": OK, "SAL": OK, "NVG Curr": OK}),
              ("Active", "S LIM JH", "A", {"Day Capt": DUE, "Day Curr": DUE, "Night Curr": "25 Oct 26", "SIM": OK, "MAS": "N/A"}),
              ("Active", "SAMY", "C", {"Day Curr": OK, "Night Curr": OK, "SIM": OK}),
              ("NSmen", "TEO WK", "A", {"Day Curr": OK, "Night Curr": "01 Oct 26"})],
    "ACS": [("Active", "L CHEW", "A", {"Day Curr": OK, "Night Curr": LAP, "SIM": OK, "MAS": DUE}), ("Active", "M SOH", "B", {"Day Curr": OK, "Night Curr": OK, "SIM": OK})],
    "TACCO": [("Active", "T LEE", "A", {"Day Curr": OK, "Night Curr": OK, "SIM": LAP, "APT": DUE}), ("Active", "R KOH", "C", {"Day Curr": OK, "Night Curr": OK, "SIM": OK})],
    "SUP1": [("Active", "P YEO", "A", {"Day Curr": OK, "Night Curr": OK, "SIM": OK}), ("Active", "Q ONG", "B", {"Day Curr": LAP, "Night Curr": OK, "SIM": OK})],
    "SUP2": [("Active", "V TAN", "A", {"Day Curr": OK, "Night Curr": "27 Oct 26", "SIM": OK, "MAS": LAP}), ("Active", "H LIM", "B", {"Day Curr": OK, "Night Curr": OK, "SIM": OK}),
             ("Staff/Assigned", "O NG", "C", {"Day Curr": DUE, "Night Curr": OK, "SIM": OK})],
}
sheets = ["PILOT", "ACS", "TACCO", "SUP1", "SUP2"]
if variant == "partial": sheets = ["PILOT", "SUP2"]

wb = openpyxl.Workbook(); wb.remove(wb.active)
for sn in sheets:
    ws = wb.create_sheet(sn)
    cols = PILOT_COLS if sn == "PILOT" else OTHER_COLS
    heads = ["Status", "Name", "CAT", "Instructor CAT"]
    ws.cell(1, 1, "Aircrew Details"); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    for i, h in enumerate(heads, 1): ws.cell(2, i, h)
    c = 5
    for g, names in cols:
        ws.cell(1, c, g)
        if len(names) > 1: ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + len(names) - 1)
        for n in names: ws.cell(2, c, n); c += 1
    last_status = None; r = 3
    for status, name, cat, vals in PEOPLE[sn]:
        if status != last_status:   # band row between status groups
            ws.cell(r, 1, " -> " + ("Active" if status == "Active" else status)); r += 1; last_status = status
        ws.cell(r, 1, status); ws.cell(r, 2, name); ws.cell(r, 3, cat)
        c = 5
        for g, names in cols:
            for n in names:
                if g in ("Flight Currencies", "SIM Currencies", "LSS Currencies", "Mission Currencies"):
                    v = vals.get(n, "-")
                    if isinstance(v, dt.datetime): cell = ws.cell(r, c, v); cell.number_format = "dd mmm yy"
                    else: ws.cell(r, c, v)
                elif n == "Remarks": ws.cell(r, c, "FAKE personal remark")
                elif n == "Level": ws.cell(r, c, "TL 3")
                else: ws.cell(r, c, 4.5)
                c += 1
        r += 1
wb.properties.created = created
wb.save(out)
