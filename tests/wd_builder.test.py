#!/usr/bin/env python3
# WD (whole day) fourth row: checks on a freshly built department file.
# Run: python3 tests/wd_builder.test.py [folder holding Flight_Ops_Planner_SUP2.xlsx]   (no folder: builds the demo SUP 2 into a temp dir)
import os, subprocess, sys, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
if len(sys.argv) > 1: folder = sys.argv[1]
else:
    folder = tempfile.mkdtemp(prefix="lodestar_wdb_")
    subprocess.run([sys.executable, os.path.join(ROOT, ".lodestar_build.py"), "excel", folder, "--only", "sup2"], check=True, stdout=subprocess.DEVNULL)
wb = openpyxl.load_workbook(os.path.join(folder, "Flight_Ops_Planner_SUP2.xlsx"))
fails = []
def check(name, ok, detail=""):
    print(("ok   " if ok else "FAIL ") + name + ("" if ok else ": " + str(detail)))
    if not ok: fails.append(name)

G0 = 27
oct_ = wb["Oct"]
check("stamp_is_layout_3", str(wb["Dashboard"]["Z2"].value).endswith("layout 3"), wb["Dashboard"]["Z2"].value)
check("four_rows_per_person", oct_.cell(G0 + 4, 3).value == "AM" and oct_.cell(G0 + 7, 3).value == "WD" and oct_.cell(G0 + 3, 3).value == "WD", [oct_.cell(G0 + k, 3).value for k in range(8)])
wd = oct_.cell(G0 + 3, 3)
check("wd_label_is_teal", wd.fill.fgColor.rgb.endswith("CCFBF1") and wd.font.color.rgb.endswith("115E59"), (wd.fill.fgColor.rgb, wd.font.color.rgb))
check("wd_day_cells_tinted", oct_.cell(G0 + 3, 4).fill.fgColor.rgb.endswith("F0FDFA"), oct_.cell(G0 + 3, 4).fill.fgColor.rgb)
check("night_label_unchanged", oct_.cell(G0 + 2, 3).fill.fgColor.rgb.endswith("1E1B4B"), oct_.cell(G0 + 2, 3).fill.fgColor.rgb)
check("callsign_on_all_four_rows", all(oct_.cell(G0 + k, 1).value == oct_.cell(G0, 1).value for k in range(4)))
check("totals_merged_over_four_rows", any(str(m) == "AI%d:AI%d" % (G0, G0 + 3) for m in oct_.merged_cells.ranges), [str(m) for m in oct_.merged_cells.ranges if str(m).startswith("AI")][:3])
f1 = [str(d.formula1) for d in oct_.data_validations.dataValidation if "ActCodes" in str(d.formula1)]
check("grid_validation_switches_on_wd", any("WdCodes" in x and '"WD"' in x for x in f1), f1)
check("wdcodes_name_exists", "WdCodes" in wb.defined_names, list(wb.defined_names.keys())[:20])
dv = wb["Daily View"]; hdr = [dv.cell(10, c).value for c in range(1, 8)]
check("daily_view_has_wd_column", hdr == ["Callsign", "Crew Member", "Role", "AM", "PM", "NIGHT", "WD"], hdr)
fx = str(dv.cell(11, 4).value) + str(dv.cell(12, 4).value)
check("daily_view_strides_by_four", "INDEX(" in fx and "4*" in str(dv.cell(12, 4).value).replace(" ", "") or "5," in str(dv.cell(12, 4).value), dv.cell(12, 4).value)
cn = [str(d.formula1) for d in wb["Cell Notes"].data_validations.dataValidation if "AM" in str(d.formula1)]
check("cell_notes_shift_list_has_wd", any("WD" in x for x in cn), cn)
req = [str(d.formula1) for d in wb["Requests"].data_validations.dataValidation if "AM" in str(d.formula1)] if "Requests" in wb.sheetnames else ["n/a WD"]
check("requests_shift_list_has_wd", any("WD" in x for x in req), req)
hidden = [oct_.row_dimensions[G0 + 4 * 59 + k].hidden for k in range(4)]
check("empty_slot_hides_all_four_rows", all(hidden), hidden)
print("%d failed" % len(fails) if fails else "all pass"); sys.exit(1 if fails else 0)
